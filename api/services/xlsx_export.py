"""Single implementation of the Excel export style contract.

Every data export in this project must go through here. The rules are recorded
in AGENTS.md under "Excel 导出样式规范（所有数据导出必须遵守）":

* real .xlsx (openpyxl) — a CSV cannot carry fonts, row heights, number formats
  or hidden gridlines
* gridlines off and no cell borders ("a table on white paper")
* Aptos Narrow 11; bold header on a light fill; header height 30, data height 24
* everything centred, except designated long-text columns which are left
  aligned with indent 1
* money / percent / date written as real numbers and dates with a number format
* column width fitted to the widest text so nothing is clipped
* frozen header row
* placeholder encodings become genuinely empty cells, never a dash
"""

from __future__ import annotations

import io
import re
from datetime import date
from typing import Any, Iterable, Sequence

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.styles.fonts import Font as _Font
from openpyxl.utils import get_column_letter

FONT_NAME = "Aptos Narrow"
FONT_SIZE = 11
TEXT_COLOR = "1F2A37"
HEADER_FILL = "F4F7FB"
HEADER_ROW_HEIGHT = 30
DATA_ROW_HEIGHT = 24
MAX_COL_WIDTH = 120
MIN_COL_WIDTH = 8

NUMBER_FORMATS = {"money": '"¥"#,##0', "percent": "0%", "percent100": "0%",
                  "date": "yyyy-mm-dd", "number": "#,##0"}

# "no value" encodings the DB still carries around
_BLANK_TOKENS = {"", "0", "00:00:00", "0000-00-00", "0000-00-00 00:00:00",
                 "null", "none", "nan", "-", "—", "--"}

# Date shapes actually seen in this project: ISO from the API, and the SPA's own
# YY/MM/DD / YY/MM output (see fmtDateG in index.html). MM/DD/YYYY is deliberately
# NOT supported — it would be ambiguous against YY/MM/DD.
_DATE_PATTERNS = (
    (re.compile(r"^(\d{4})-(\d{2})-(\d{2})"), "ymd"),
    (re.compile(r"^(\d{4})-(\d{2})$"),          "ym"),
    (re.compile(r"^(\d{4})/(\d{2})/(\d{2})"),  "ymd"),
    (re.compile(r"^(\d{2})/(\d{2})/(\d{2})$"), "ymd2"),
    (re.compile(r"^(\d{2})/(\d{2})$"),          "ym2"),
)


def _to_date(text: str):
    """Parse a date from the shapes this app produces, else None.

    Returns ``"blank"`` for the far-future "open ended" sentinel (year >= 9000).
    """
    for pattern, shape in _DATE_PATTERNS:
        m = pattern.match(text)
        if not m:
            continue
        g = m.groups()
        if shape == "ym":
            y, mo, d = int(g[0]), int(g[1]), 1
        elif shape == "ym2":
            y, mo, d = 2000 + int(g[0]), int(g[1]), 1
        elif shape == "ymd2":
            y, mo, d = 2000 + int(g[0]), int(g[1]), int(g[2])
        else:
            y, mo, d = int(g[0]), int(g[1]), int(g[2])
        if y >= 9000:
            return "blank"
        try:
            return date(y, mo, d)
        except ValueError:
            return None
    return None


def _to_number(text: str):
    cleaned = re.sub(r"[^0-9.\-]", "", text)
    if cleaned in ("", "-", ".", "-."):
        return None
    try:
        return float(cleaned)
    except ValueError:
        return None


def normalise(kind: str, raw: Any):
    """Return ``(excel_value, display_text)`` for one cell.

    ``kind`` is one of ``text`` / ``number`` / ``money`` / ``percent`` /
    ``percent100`` / ``date``. Use ``percent100`` when the caller knows the value
    is already on a 0-100 scale (YoY deltas); ``percent`` is the tolerant one.
    Every numeric kind is returned as a real number so Excel shows it formatted
    and can sum and sort it. Money keeps its currency format, percent is stored
    as a fraction, dates become real dates.
    """
    if raw is None:
        return None, ""
    text = str(raw).strip()
    if text == "":
        return None, ""

    # Numeric kinds first: a real 0 is data, not an empty cell. Only the
    # string-encoded placeholders of text/date columns mean "no value".
    if kind in ("number", "money", "percent", "percent100"):
        # a dash placeholder means "no value" even in a numeric column, but a
        # real 0 is data and must survive
        if text in ("-", "—", "--", "n/a", "null", "none", "nan"):
            return None, ""
        num = _to_number(text)
        if num is None:
            return text, text
        if kind == "money":
            return num, f"¥{round(num):,}"
        if kind == "percent":
            if abs(num) > 1.5:          # stored as 25 rather than 0.25
                num = num / 100.0
            return num, f"{round(num * 100)}%"
        if kind == "percent100":
            # explicitly on a 0-100 scale (e.g. a YoY delta of 12.5) — never guess
            return num / 100.0, f"{round(num)}%"
        if float(num).is_integer():
            num = int(num)
        return num, f"{round(num):,}"

    if text.lower() in _BLANK_TOKENS:
        return None, ""

    if kind == "date":
        parsed = _to_date(text)
        if parsed == "blank":
            return None, ""
        if parsed is not None:
            return parsed, parsed.isoformat()
        return text, text

    return text, text


def render_sheet(ws, columns: Sequence[dict], rows: Iterable[Sequence[Any]]) -> None:
    """Write one homogeneous table onto ``ws`` following the style contract.

    ``columns`` items: ``{"label": str, "kind": str, "align": "center"|"left",
    "width_min": int}``. ``rows`` items are sequences aligned with ``columns``.
    """
    header_font = Font(name=FONT_NAME, bold=True, size=FONT_SIZE, color=TEXT_COLOR)
    data_font = Font(name=FONT_NAME, size=FONT_SIZE, color=TEXT_COLOR)
    centre = Alignment(horizontal="center", vertical="center")
    header_fill = PatternFill("solid", fgColor=HEADER_FILL)

    def align_for(col: dict):
        if str(col.get("align", "center")).lower() == "left":
            # long list-like text (e.g. "A+B+C") reads better left aligned + indent
            return Alignment(horizontal="left", vertical="center", indent=1)
        return centre

    widths = []
    for ci, col in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=ci, value=col.get("label", ""))
        cell.font, cell.alignment, cell.fill = header_font, centre, header_fill
        widths.append(max(int(col.get("width_min") or MIN_COL_WIDTH), len(str(col.get("label", "")))))
    ws.row_dimensions[1].height = HEADER_ROW_HEIGHT

    ri = 1
    for row in rows:
        ri += 1
        for ci, col in enumerate(columns, start=1):
            raw = row[ci - 1] if ci - 1 < len(row) else None
            value, display = normalise(str(col.get("kind") or "text"), raw)
            cell = ws.cell(row=ri, column=ci, value=value)
            cell.font, cell.alignment = data_font, align_for(col)
            numfmt = NUMBER_FORMATS.get(str(col.get("kind") or "text"))
            if numfmt and value is not None:
                cell.number_format = numfmt
            if len(display) > widths[ci - 1]:
                widths[ci - 1] = len(display)
        ws.row_dimensions[ri].height = DATA_ROW_HEIGHT

    for ci, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(ci)].width = min(width + 3, MAX_COL_WIDTH)


def build_workbook(sheet_name: str, columns: Sequence[dict], rows: Iterable[Sequence[Any]]):
    """Return an in-memory .xlsx as a BytesIO, styled per the contract."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = (sheet_name or "Export")[:31]
    ws.sheet_view.showGridLines = False
    render_sheet(ws, columns, rows)
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _restyle_font(font):
    """Same font, different typeface — keeps size / bold / colour / italic."""
    return _Font(name=FONT_NAME, size=font.size, bold=font.bold, italic=font.italic,
                 color=font.color, underline=font.underline, strike=font.strike)


def normalise_workbook(wb, row_height_floor: int | None = None):
    """Bring an already-built workbook onto the style contract.

    Used by the complex reports that keep their own hand-designed layout
    (merged headers, section colours, pivots): their structure stays, but they
    must still be gridline-free and share the house typeface. Returns the number
    of cells whose font was changed.
    """
    changed = 0
    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = False
        for row in ws.iter_rows():
            for cell in row:
                if cell.font and cell.font.name != FONT_NAME:
                    cell.font = _restyle_font(cell.font)
                    changed += 1
            if row_height_floor:
                dim = ws.row_dimensions[row[0].row]
                if dim.height is None or dim.height < row_height_floor:
                    dim.height = row_height_floor
    return changed


XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"