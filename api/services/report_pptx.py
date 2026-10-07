"""Deterministic editable PPTX export for fixed 1920×1080 Report Decks.

Chromium measures the rendered DOM; this writer creates native PowerPoint
objects. Unsupported content fails rather than becoming a slide screenshot.
"""
from __future__ import annotations

import base64
import io
import json
import re
from pathlib import Path
from urllib.parse import urlparse

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

from services.report_pptx_fonts import embed_project_fonts

CANVAS_W, CANVAS_H = 1920, 1080
PX_PER_IN = 144
PPTX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
EXTRACTOR = Path(__file__).with_name("report_pptx_extract.js")


class UnsupportedReport(ValueError):
    """A report cannot meet the native-object export contract."""


def _in(px):
    return Inches(float(px) / PX_PER_IN)


def _pt(px):
    return Pt(max(1, float(px) / 2))


def _stroke_pt(px):
    # A 1 px CSS edge is 0.5 pt on the 144 px/in export canvas. The text-size
    # helper has a 1 pt floor, which made every fine badge border twice as thick.
    return Pt(float(px) / 2)


def _color(value):
    if not value:
        return None
    value = value.strip().lower()
    if value.startswith("#"):
        raw = value[1:]
        if len(raw) == 3:
            raw = "".join(c * 2 for c in raw)
        if len(raw) == 6:
            return RGBColor.from_string(raw.upper())
    if value.startswith(("rgb(", "rgba(")):
        bits = value[value.index("(") + 1:value.index(")")].split(",")
        if len(bits) >= 4 and float(bits[3].strip()) == 0:
            return None
        return RGBColor(*(max(0, min(255, round(float(v)))) for v in bits[:3]))
    return None


def _box(obj):
    return tuple(_in(obj[k]) for k in ("x", "y", "w", "h"))


def _font(font, style, text=""):
    family = (style.get("font") or "Arial").split(",")[0].strip(" ' \"") or "Arial"
    if re.search(r"[\u3400-\u9fff]", text) and "CoolSans SC Narrow" in (style.get("font") or ""):
        family = "CoolSans SC Narrow"
    font.name = family
    font.size = _pt(style.get("size") or 18)
    font.bold = bool(style.get("bold"))
    font.italic = bool(style.get("italic"))
    font.underline = bool(style.get("underline"))
    color = _color(style.get("color"))
    if color:
        font.color.rgb = color


def _write_text(shape, obj):
    paragraphs = obj.get("paragraphs") or [{
        "style": obj.get("style") or {},
        "runs": [{"text": obj.get("text") or "", "style": obj.get("style") or {}}],
    }]
    if not any(run.get("text") for para in paragraphs for run in para["runs"]):
        return
    frame = shape.text_frame
    frame.clear()
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    margins = obj.get("margins") or {}
    for side in ("left", "right", "top", "bottom"):
        if side in margins:
            setattr(frame, "margin_" + side, _in(margins[side]))
    frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE if obj.get("fit") else MSO_AUTO_SIZE.NONE
    frame.word_wrap = bool(obj.get("wrap", False))
    if obj.get("valign") == "middle":
        frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    if obj.get("padding"):
        frame.margin_left = frame.margin_right = _in(obj["padding"])
    for index, block in enumerate(paragraphs):
        para = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        style = block.get("style") or {}
        para.alignment = {"center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}.get(style.get("align"), PP_ALIGN.LEFT)
        para.space_before = Pt(0)
        para.space_after = Pt(float(block.get("spaceAfter") or 0) / 2)
        if style.get("lineHeight"):
            para.line_spacing = _pt(style["lineHeight"])
        for fragment in block["runs"]:
            run_style = fragment.get("style") or style
            # DOM <br> is emitted as a vertical tab by the extractor.  Passing
            # that character to python-pptx escapes it as the visible string
            # "_x000B_"; add a real DrawingML line break instead.
            for line_index, line_text in enumerate((fragment.get("text") or "").split("\v")):
                if line_index:
                    para.add_line_break()
                for part in re.findall(r"[\u3400-\u9fff]+|[^\u3400-\u9fff]+", line_text):
                    run = para.add_run()
                    run.text = part
                    _font(run.font, run_style, part)


def _text(slide, obj):
    x, y, w, h = _box(obj)
    if w <= 0 or h <= 0:
        return
    _write_text(slide.shapes.add_textbox(x, y, w, h), obj)


def _remove_theme_effect(shape):
    # python-pptx writes a p:style with effectRef idx=1 for both auto-shapes
    # and connectors. Even an explicitly coloured connector can inherit an
    # Office theme shadow through that reference.
    for theme_style in shape._element.xpath("./p:style"):
        shape._element.remove(theme_style)


def _shape(slide, obj):
    x, y, w, h = _box(obj)
    if w <= 0 or h <= 0:
        return
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if obj.get("radius", 0) >= 2 else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, x, y, w, h)
    _remove_theme_effect(shape)
    if kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        shape.adjustments[0] = min(0.5, float(obj["radius"]) / max(1, min(float(obj["w"]), float(obj["h"]))))
    shape.line.fill.background()
    fill = _color(obj.get("fill"))
    if fill:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
    else:
        shape.fill.background()
    stroke = obj.get("stroke") or {}
    stroke_color = _color(stroke.get("color"))
    if stroke_color and stroke.get("width", 0) > 0:
        shape.line.color.rgb = stroke_color
        shape.line.width = _stroke_pt(stroke["width"])
    # A crisp, opaque CSS ring is an outline, not a PowerPoint shadow. Blurred
    # shadows are deliberately omitted: Office renders their alpha differently
    # and can make nearly invisible HTML decoration look like a dark halo.
    shadow = obj.get("shadow") or {}
    if (shadow and not shadow.get("x") and not shadow.get("y") and
            not shadow.get("blur") and shadow.get("spread", 0) > 0):
        color = _color(shadow.get("color"))
        if color:
            shape.line.color.rgb = color
            shape.line.width = _stroke_pt(shadow["spread"] * 2)
    for side in ("top", "right", "bottom", "left"):
        border = (obj.get("borders") or {}).get(side) or {}
        color = _color(border.get("color"))
        if color and border.get("width", 0) > 0:
            # CSS borders are painted *inside* the measured border box.
            # Centering an Office connector on the outer edge instead makes
            # adjacent grid lines look darker and thicker than the HTML.
            half = _in(float(border["width"]) / 2)
            coords = {
                "top": (x, y + half, x + w, y + half),
                "right": (x + w - half, y, x + w - half, y + h),
                "bottom": (x, y + h - half, x + w, y + h - half),
                "left": (x + half, y, x + half, y + h),
            }[side]
            line = slide.shapes.add_connector(1, *coords)
            _remove_theme_effect(line)
            line.line.color.rgb = color
            line.line.width = _stroke_pt(border["width"])
    return shape


def _badge(slide, obj):
    """A badge is one editable PowerPoint shape, including its text and edge."""
    shape = _shape(slide, obj)
    if shape is not None:
        # A badge's measured bounds already include its CSS padding.  Native
        # text must keep its authored size when a slide is pasted into another
        # deck, even if that deck substitutes a font and recalculates layout.
        text = {**obj, "fit": False, "margins": dict.fromkeys(("left", "right", "top", "bottom"), 0)}
        text.pop("padding", None)
        _write_text(shape, text)


def _table_cell_borders(cell, borders):
    """Write browser-measured CSS edges and suppress Office's theme grid."""
    tc_pr = cell._tc.get_or_add_tcPr()
    for index, (side, tag) in enumerate((("left", "a:lnL"), ("right", "a:lnR"),
                                         ("top", "a:lnT"), ("bottom", "a:lnB"))):
        border = borders.get(side) or {}
        line = OxmlElement(tag)
        color = _color(border.get("color"))
        if color and border.get("width", 0) > 0:
            line.set("w", str(int(_stroke_pt(border["width"]))))
            fill = OxmlElement("a:solidFill")
            rgb = OxmlElement("a:srgbClr")
            rgb.set("val", str(color))
            fill.append(rgb)
            line.append(fill)
        else:
            line.append(OxmlElement("a:noFill"))
        tc_pr.insert(index, line)


def _table(slide, obj):
    rows = obj["rows"]
    if not rows or not rows[0]:
        return
    occupied = set()
    placements = []
    widths = {}
    for i, row in enumerate(rows):
        col = 0
        for item in row:
            while (i, col) in occupied:
                col += 1
            rs, cs = int(item.get("rowspan", 1)), int(item.get("colspan", 1))
            if rs < 1 or cs < 1 or i + rs > len(rows):
                raise UnsupportedReport("HTML 表格 span 不合法 / invalid HTML table span")
            if any((r, c) in occupied for r in range(i, i + rs) for c in range(col, col + cs)):
                raise UnsupportedReport("HTML 表格单元格重叠 / overlapping HTML table cells")
            for r in range(i, i + rs):
                for c in range(col, col + cs):
                    occupied.add((r, c))
            placements.append((i, col, rs, cs, item))
            if cs == 1:
                widths.setdefault(col, float(item["w"]))
            col += cs
    cols = max(c for _, c in occupied) + 1
    if any((r, c) not in occupied for r in range(len(rows)) for c in range(cols)):
        raise UnsupportedReport("HTML 表格网格不规则 / irregular HTML table grid")
    remaining = max(1, float(obj["w"]) - sum(widths.values()))
    missing = cols - len(widths)
    for c in range(cols):
        widths.setdefault(c, remaining / missing if missing else float(obj["w"]) / cols)
    table = slide.shapes.add_table(len(rows), cols, *_box(obj)).table
    # python-pptx attaches PowerPoint's default table style, which adds fills,
    # grid lines and accents not present in the report. Every visible cell style
    # must come from the measured HTML instead.
    table_props = table._tbl.tblPr
    for theme_style in table_props.xpath("./a:tableStyleId"):
        table_props.remove(theme_style)
    for flag in ("firstRow", "firstCol", "lastRow", "lastCol", "bandRow", "bandCol"):
        table_props.set(flag, "0")
    for c in range(cols):
        table.columns[c].width = _in(widths[c])
    row_heights = obj.get("rowHeights") or []
    for i, row in enumerate(rows):
        height = row_heights[i] if i < len(row_heights) else max(
            item["h"] / item.get("rowspan", 1) for item in row
        )
        table.rows[i].height = _in(height)
    for i, j, rs, cs, item in placements:
        cell = table.cell(i, j)
        if rs > 1 or cs > 1:
            cell.merge(table.cell(i + rs - 1, j + cs - 1))
        cell.text = item["text"]
        padding = item.get("padding") or {}
        cell.margin_left = _in(padding.get("left", 10))
        cell.margin_right = _in(padding.get("right", 10))
        cell.margin_top = _in(padding.get("top", 4))
        cell.margin_bottom = _in(padding.get("bottom", 4))
        cell.vertical_anchor = {"top": MSO_ANCHOR.TOP, "bottom": MSO_ANCHOR.BOTTOM}.get(
            item.get("verticalAlign"), MSO_ANCHOR.MIDDLE)
        style = item["style"]
        bg = _color(style.get("background"))
        if bg:
            cell.fill.solid()
            cell.fill.fore_color.rgb = bg
        else:
            cell.fill.background()
        _table_cell_borders(cell, item.get("borders") or {})
        for para in cell.text_frame.paragraphs:
            para.alignment = {"center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}.get(style.get("align"), PP_ALIGN.LEFT)
            para.space_before = para.space_after = Pt(0)
            for run in para.runs:
                parts = re.findall(r"[\u3400-\u9fff]+|[^\u3400-\u9fff]+", run.text)
                if len(parts) <= 1:
                    _font(run.font, style, run.text)
                else:
                    run.text = parts[0]
                    _font(run.font, style, parts[0])
                for part in parts[1:]:
                    extra = para.add_run()
                    extra.text = part
                    _font(extra.font, style, part)


def _line(slide, obj):
    color = _color(obj.get("color"))
    if color is None:
        return
    line = slide.shapes.add_connector(1, _in(obj["x"]), _in(obj["y"]),
                                      _in(obj["x"] + obj["w"]), _in(obj["y"] + obj["h"]))
    _remove_theme_effect(line)
    line.line.color.rgb = color
    line.line.width = _stroke_pt(obj.get("width") or 1)


def _chart(slide, obj):
    data = obj["data"]
    kind = {
        "bar": XL_CHART_TYPE.BAR_CLUSTERED, "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
        "line": XL_CHART_TYPE.LINE, "pie": XL_CHART_TYPE.PIE,
    }.get(data.get("type"))
    if kind is None:
        raise UnsupportedReport("不支持的原生图表类型 / unsupported native chart type")
    categories, series = data.get("categories") or [], data.get("series") or []
    if not categories or not series or any(len(s.get("values") or []) != len(categories) for s in series):
        raise UnsupportedReport("原生图表要求类目与数值序列等长 / native chart requires categories and numeric series of equal length")
    if any(not isinstance(v, (int, float)) for s in series for v in s["values"]):
        raise UnsupportedReport("原生图表的值必须是数字 / native chart values must be numeric")
    chart_data = CategoryChartData()
    chart_data.categories = categories
    for item in series:
        chart_data.add_series(item.get("name") or "Series", item["values"])
    chart = slide.shapes.add_chart(kind, *_box(obj), chart_data).chart
    chart.has_legend = bool(data.get("legend", len(series) > 1))
    for index, item in enumerate(series):
        color = _color(item.get("color"))
        if color:
            chart.series[index].format.fill.solid()
            chart.series[index].format.fill.fore_color.rgb = color
            chart.series[index].format.line.color.rgb = color


def _image(slide, obj):
    data = obj.get("data") or ""
    if not data.startswith("data:image/") or ";base64," not in data:
        raise UnsupportedReport("图片无法嵌入；请用站内地址或 data URL / image could not be embedded; use an in-app or data URL")
    raw = base64.b64decode(data.split(";base64,", 1)[1], validate=True)
    if len(raw) > 12 * 1024 * 1024:
        raise UnsupportedReport("有图片超过 12 MB 导出上限 / an image exceeds the 12 MB export limit")
    slide.shapes.add_picture(io.BytesIO(raw), *_box(obj))


def build_pptx(scene):
    """Turn a measured scene into native text, shapes, tables, charts and images."""
    pages = scene.get("slides") or []
    if not pages:
        raise UnsupportedReport("只有带 .slide 页的 16:9 Report Deck 才能导出 / only 16:9 Report Decks with .slide pages can be exported")
    if len(pages) > 100:
        raise UnsupportedReport("PPTX 导出最多 100 页 / PPTX export is limited to 100 slides")
    pres = Presentation()
    pres.slide_width, pres.slide_height = Inches(40 / 3), Inches(7.5)
    for page in pages:
        nodes = page.get("nodes") or []
        if len(nodes) > 8000:
            raise UnsupportedReport("某一页的可编辑对象过多 / a slide has too many editable objects")
        slide = pres.slides.add_slide(pres.slide_layouts[6])
        bg = _color(page.get("background"))
        if bg:
            slide.background.fill.solid()
            slide.background.fill.fore_color.rgb = bg
        for obj in nodes:
            kind = obj.get("kind")
            handler = {"text": _text, "shape": _shape, "badge": _badge, "table": _table,
                       "chart": _chart, "image": _image, "line": _line}.get(kind)
            if handler is None:
                raise UnsupportedReport(f"不支持的页面对象: {kind} / unsupported slide object: {kind}")
            handler(slide, obj)
    output = io.BytesIO()
    pres.save(output)
    # The browser uses the bundled web faces, but a font name alone leaves
    # PowerPoint free to substitute an installed face.  Embed the corresponding
    # full editable faces when this scene actually uses them.
    scene_text = json.dumps(scene, ensure_ascii=False)
    families = {name for name in ("Roboto Condensed", "CoolSans SC Narrow") if name in scene_text}
    return embed_project_fonts(output.getvalue(), families)


async def render_report_to_pptx(report_url, lang="", state_json="{}", data_snapshots=None,
                                html_body=None, asset_auth_headers=None):
    """Render the selected report snapshot in isolated Chromium before building its PPTX."""
    from playwright.async_api import async_playwright

    parsed = urlparse(report_url)
    if parsed.hostname not in {"127.0.0.1", "localhost"} or parsed.port != 8000:
        raise UnsupportedReport("导出器只接受本地报告路由 / the exporter accepts only the local report route")
    # Only the report's own read-only state endpoint is available to its script.
    # Serve the same DB snapshot selected with the HTML instead of permitting a
    # second, racing request to the API. In particular, never permit POST here.
    report_path = parsed.path.rstrip("/")
    if "/r/" not in report_path:
        raise UnsupportedReport("导出器只接受本地报告路由 / the exporter accepts only the local report route")
    base_path, slug = report_path.rsplit("/r/", 1)
    state_path = f"{base_path}/api/reports/{slug}/state"
    snapshots = data_snapshots or {}
    async with async_playwright() as playwright:
        browser = await playwright.chromium.launch(headless=True, args=["--no-sandbox"])
        try:
            page = await browser.new_page(viewport={"width": CANVAS_W, "height": CANVAS_H}, device_scale_factor=1)
            missing_data = []

            async def restrict(route):
                url = urlparse(route.request.url)
                same = url.hostname in {"127.0.0.1", "localhost"} and url.port == 8000
                asset = url.path.lower().endswith((".png", ".jpg", ".jpeg", ".webp", ".gif", ".woff", ".woff2", ".css"))
                if same and url.path == state_path and route.request.method == "GET":
                    await route.fulfill(status=200, content_type="application/json",
                                        headers={"Cache-Control": "no-store"}, body=state_json or "{}")
                elif same and route.request.method == "GET" and (key := url.path + ("?" + url.query if url.query else "")) in snapshots:
                    await route.fulfill(status=200, content_type="application/json", body=snapshots[key])
                elif same and route.request.method == "GET" and url.path == parsed.path and html_body is not None:
                    # The report route now requires a logged-in reader. Serve the
                    # already-authorized DB snapshot, without handing a session
                    # cookie to report-authored JavaScript in Chromium.
                    await route.fulfill(status=200, content_type="text/html; charset=utf-8",
                                        body=html_body)
                elif same and route.request.method == "GET" and (url.path == parsed.path or asset):
                    headers = dict(route.request.headers)
                    headers.update(asset_auth_headers or {})
                    await route.continue_(headers=headers)
                else:
                    if same and url.path.startswith(f"{base_path}/api/"):
                        missing_data.append(url.path)
                    await route.abort()

            await page.route("**/*", restrict)
            response = await page.goto(report_url, wait_until="domcontentloaded", timeout=30000)
            if response is None or response.status != 200:
                raise UnsupportedReport("报告页面加载失败，无法导出 / report page could not be loaded for export")
            await page.wait_for_function("document.documentElement.classList.contains('deck-ready')", timeout=10000)
            await page.evaluate("document.fonts.ready")
            # Interactive reports hydrate the DOM asynchronously after fetching
            # /state. deck-ready only means the page layout is installed, so wait
            # for the state response and its subsequent render before measuring.
            await page.wait_for_load_state("networkidle", timeout=10000)
            if missing_data:
                raise UnsupportedReport("报告数据不可用，无法导出: " + ", ".join(sorted(set(missing_data))[:3]) + " / report data was not available for export: " + ", ".join(sorted(set(missing_data))[:3]))
            source = EXTRACTOR.read_text(encoding="utf-8")
            scene = await page.evaluate("(" + source + ")(" + json.dumps(lang) + ")")
            if scene.get("error"):
                raise UnsupportedReport(scene["error"])
            # CSS material effects (layered gradients/textures) have no faithful
            # native PowerPoint equivalent. Capture only their painted backdrop;
            # nested labels are emitted separately as editable text and shapes.
            for slide in scene.get("slides") or []:
                for obj in slide.get("nodes") or []:
                    if obj.get("kind") != "raster":
                        continue
                    selector = f'[data-pptx-raster-id="{obj["id"]}"]'
                    await page.evaluate("""({selector, includeChildren}) => {
                      const el=document.querySelector(selector);
                      window.__pptxHidden=includeChildren ? [] : [...el.querySelectorAll('*')].map(child=>[child,child.style.visibility]);
                      for (const [child] of window.__pptxHidden) child.style.visibility='hidden';
                    }""", {"selector": selector, "includeChildren": bool(obj.get("includeChildren"))})
                    try:
                        png = await page.locator(selector).screenshot(animations="disabled")
                    finally:
                        await page.evaluate("""() => {
                          for (const [child,old] of window.__pptxHidden||[]) child.style.visibility=old;
                          window.__pptxHidden=[];
                        }""")
                    obj["kind"] = "image"
                    obj["data"] = "data:image/png;base64," + base64.b64encode(png).decode("ascii")
            return build_pptx(scene)
        finally:
            await browser.close()
