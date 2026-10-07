"""Document cards (kind='document'): the pptx / xlsx / pdf / docx upload half of Workspace.

Pure functions plus the reader page and the page counter — no database, no HTTP, no OSS.
The transport contract is the part most worth pinning down: an agent holds an API
credential, not a storage one, so bytes arrive as base64 inside JSON (the same convention a
knowledge-page image uses) and the SERVER decides whether they are what the name claims.
Every rejection here is a case where the alternative is a reader being handed a 24 MB file
that opens as garbage.
"""
import asyncio
import base64
import io
import unittest
from datetime import datetime, timezone
from unittest.mock import patch

from fastapi import HTTPException

from routers import reports


def _pptx_bytes(slides: int = 2) -> bytes:
    from pptx import Presentation
    presentation = Presentation()
    for index in range(slides):
        slide = presentation.slides.add_slide(presentation.slide_layouts[5])
        slide.shapes.title.text = f"Slide {index + 1}"
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def _pdf_bytes(pages: int = 2) -> bytes:
    from services.pdf_io import text_pdf
    return text_pdf([[(72, 72, 12, f"page {index + 1}")] for index in range(pages)])


def _xlsx_bytes() -> bytes:
    from openpyxl import Workbook
    book = Workbook()
    first = book.active
    first.title = "Q3"
    first.append(["channel", "gto"])
    first.append(["online", 1200])
    second = book.create_sheet("Hidden")
    second.append(["x"])
    second.sheet_state = "hidden"
    fourth = book.create_sheet("Q4")
    fourth.append(["channel", "spto"])
    fourth.append(["offline", 300])
    book.create_sheet("Empty")
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


class DocTypeTests(unittest.TestCase):
    def test_only_the_four_types_are_accepted(self):
        for name in ("deck.pptx", "book.xlsx", "report.pdf", "brief.docx"):
            self.assertEqual(reports._doc_type_of(name)[0], name.rsplit(".", 1)[1])
        for name in ("old.ppt", "old.doc", "old.xls", "data.csv", "archive.zip", "noext"):
            with self.assertRaises(HTTPException) as ctx:
                reports._doc_type_of(name)
            self.assertEqual(ctx.exception.status_code, 400)

    def test_the_legacy_formats_name_the_fix_instead_of_the_problem(self):
        with self.assertRaises(HTTPException) as ctx:
            reports._doc_type_of("Q3 review.ppt")
        self.assertIn(".pptx", str(ctx.exception.detail))

    def test_a_directory_is_stripped_off_the_name(self):
        _, name = reports._doc_type_of("../../etc/passwd.pptx")
        self.assertEqual(name, "passwd.pptx")
        _, name = reports._doc_type_of(r"C:\Users\example\Deck.PPTX")
        self.assertEqual(name, "Deck.PPTX")
        self.assertEqual(reports._doc_type_of("x.pptx")[0], "pptx")

    def test_the_mime_type_is_the_fallback_when_the_name_has_no_extension(self):
        doc_type, name = reports._doc_type_of("deck", reports.DOC_MIME["pptx"])
        self.assertEqual(doc_type, "pptx")
        self.assertEqual(name, "deck")

    def test_a_control_character_in_the_name_cannot_reach_the_header(self):
        _, name = reports._doc_type_of("bad\x00\r\nname.pptx")
        self.assertEqual(name, "badname.pptx")

    def test_the_title_comes_from_the_filename_when_none_is_sent(self):
        # Underscores become spaces (they are almost always stand-ins for one); hyphens are
        # left alone — they carry meaning in real file names (`2026-Q3-review`, `G-to-G`),
        # and mangling a name is worse than an imperfect title (the agent can send `title`).
        self.assertEqual(reports._title_from_filename("Q3_channel review.pptx"), "Q3 channel review")
        self.assertEqual(reports._title_from_filename("Q3_channel-review.pptx"), "Q3 channel-review")
        self.assertEqual(reports._title_from_filename("deck.pptx"), "deck")
        self.assertEqual(reports._title_from_filename(""), "Untitled document")
        self.assertEqual(len(reports._title_from_filename("x" * 400 + ".pptx")), reports.MAX_TITLE_LEN)


class DecodeDocumentTests(unittest.TestCase):
    def _body(self, filename: str, data: bytes, **kwargs):
        return reports.DocumentIn(filename=filename,
                                  content_base64=base64.b64encode(data).decode("ascii"), **kwargs)

    def test_a_real_deck_survives_the_round_trip(self):
        raw = _pptx_bytes(2)
        doc_type, name, data = reports._decode_document(self._body("deck.pptx", raw))
        self.assertEqual(doc_type, "pptx")
        self.assertEqual(name, "deck.pptx")
        self.assertEqual(data, raw)

    def test_a_data_uri_is_accepted_and_its_declared_type_is_ignored(self):
        raw = _pptx_bytes(1)
        body = reports.DocumentIn(
            filename="deck.pptx",
            content_base64="data:application/zip;base64," + base64.b64encode(raw).decode("ascii"))
        self.assertEqual(reports._decode_document(body)[0], "pptx")

    def test_bytes_that_are_not_what_the_name_claims_are_rejected(self):
        with self.assertRaises(HTTPException) as ctx:
            reports._decode_document(self._body("deck.pptx", b"MZ\x90\x00 not a zip"))
        self.assertIn("OOXML", str(ctx.exception.detail))
        with self.assertRaises(HTTPException) as ctx:
            reports._decode_document(self._body("report.pdf", b"PK\x03\x04 a zip"))
        self.assertIn("%PDF", str(ctx.exception.detail))

    def test_empty_and_malformed_payloads_are_rejected(self):
        with self.assertRaises(HTTPException):
            reports._decode_document(reports.DocumentIn(filename="deck.pptx", content_base64=""))
        with self.assertRaises(HTTPException):
            reports._decode_document(reports.DocumentIn(filename="deck.pptx", content_base64="not base64!!"))
        with self.assertRaises(HTTPException) as ctx:
            reports._decode_document(reports.DocumentIn(
                filename="deck.pptx", content_base64=base64.b64encode(b"").decode("ascii")))
        self.assertIn("empty", str(ctx.exception.detail))

    def test_the_size_limit_is_enforced_with_the_numbers_in_the_message(self):
        original = reports.MAX_DOC_BYTES
        reports.MAX_DOC_BYTES = 16
        try:
            with self.assertRaises(HTTPException) as ctx:
                reports._decode_document(self._body("deck.pptx", _pptx_bytes(1)))
            self.assertEqual(ctx.exception.status_code, 413)
            self.assertIn("16", str(ctx.exception.detail))
        finally:
            reports.MAX_DOC_BYTES = original


class DocumentMetadataTests(unittest.TestCase):
    def test_page_counts_come_from_the_real_file(self):
        self.assertEqual(reports._document_pages("pptx", _pptx_bytes(3)), 3)
        self.assertEqual(reports._document_pages("pdf", _pdf_bytes(2)), 2)
        # Three VISIBLE sheets (Q3 / Q4 / Empty); the hidden one is not a "page".
        self.assertEqual(reports._document_pages("xlsx", _xlsx_bytes()), 3)
        self.assertIsNone(reports._document_pages("docx", b"PK\x03\x04 not really a docx"))

    def test_human_size(self):
        self.assertEqual(reports._human_size(0), "0 B")
        self.assertEqual(reports._human_size(999), "999 B")
        self.assertEqual(reports._human_size(2048), "2.0 KB")
        self.assertEqual(reports._human_size(24 * 1024 * 1024), "24.0 MB")

    def test_a_pdf_opens_inline_and_everything_else_downloads(self):
        pdf = {"slug": "s", "doc_type": "pdf", "doc_mime": reports.DOC_MIME["pdf"], "doc_name": "r.pdf"}
        deck = {"slug": "s", "doc_type": "pptx", "doc_mime": reports.DOC_MIME["pptx"], "doc_name": "d.pptx"}
        self.assertTrue(reports._document_headers(pdf, download=False)["Content-Disposition"].startswith("inline"))
        self.assertTrue(reports._document_headers(pdf, download=True)["Content-Disposition"].startswith("attachment"))
        self.assertTrue(reports._document_headers(deck, download=False)["Content-Disposition"].startswith("attachment"))

    def test_a_non_ascii_filename_still_has_an_ascii_fallback(self):
        row = {"slug": "s", "doc_type": "pptx", "doc_name": "季度复盘.pptx"}
        disposition = reports._document_headers(row, download=True)["Content-Disposition"]
        self.assertIn("filename=", disposition)
        self.assertIn("filename*=UTF-8''", disposition)
        # The ASCII part must not be empty (some clients take it literally).
        ascii_part = disposition.split('filename="', 1)[1].split('"', 1)[0]
        self.assertTrue(ascii_part)


class DocumentPageTests(unittest.TestCase):
    def _row(self, **kwargs):
        row = {"slug": "q3-deck", "title": "Q3 deck", "doc_type": "pptx",
               "doc_name": "Q3_Deck.pptx", "size_bytes": 2048, "doc_pages": 12}
        row.update(kwargs)
        return row

    def test_a_pptx_page_frames_the_html_preview_and_offers_the_original(self):
        # ⚠️ The preview is HTML, not a server-made PDF: it is rendered by the READER'S browser,
        # whose fonts include Chinese — the API image has no CJK font at all, and the PDF version
        # showed boxes for every Chinese character (2026-10-01).
        html = reports._document_page_html(self._row(), file_url="api/reports/q3-deck/document",
                                           preview_url="api/reports/q3-deck/document/preview.html")
        self.assertIn('src="api/reports/q3-deck/document/preview.html"', html)
        self.assertIn("api/reports/q3-deck/document?download=1", html)
        self.assertIn("Download", html)
        self.assertNotIn("<script", html)          # the reader page is inert on purpose

    def test_a_pdf_page_frames_the_html_preview_too(self):
        # ⚠️ Not the file itself any more (2026-10-01): a PDF is rasterised into sheets of paper
        # like a docx, because a note layer cannot be drawn over a browser PDF plugin. The reader
        # page therefore frames the same preview route as the other three types.
        html = reports._document_page_html(self._row(doc_type="pdf"), file_url="api/reports/x/document",
                                           preview_url="api/reports/x/document/preview.html")
        self.assertIn('src="api/reports/x/document/preview.html"', html)
        self.assertIn("api/reports/x/document?download=1", html)

    def test_an_xlsx_page_frames_the_grid_preview(self):
        html = reports._document_page_html(self._row(doc_type="xlsx"),
                                           file_url="api/reports/x/document",
                                           preview_url="api/reports/x/document/preview.html")
        self.assertIn('src="api/reports/x/document/preview.html"', html)

    def test_one_page_for_all_four_types(self):
        # The reader page no longer has a per-type body: whatever the file is, the page frames
        # the preview and offers the download. A second rendering path here is how the xlsx page
        # drifted from the pptx one in the first place.
        for doc_type in ("pptx", "docx", "xlsx", "pdf"):
            html = reports._document_page_html(self._row(doc_type=doc_type), file_url="api/reports/x/document",
                                               preview_url="api/reports/x/document/preview.html")
            self.assertIn('src="api/reports/x/document/preview.html"', html, doc_type)
            self.assertEqual(html.count("<iframe"), 1, doc_type)
            self.assertIn("download=1", html, doc_type)

    def test_the_reader_page_says_notes_are_available(self):
        html = reports._document_page_html(self._row(), file_url="u", preview_url="p")
        self.assertIn("Right-click", html)

    def test_the_title_is_escaped_and_the_guest_view_is_labeled(self):
        html = reports._document_page_html(self._row(title='<img src=x onerror="boom()">'),
                                           file_url="s/tok/document", preview_url="s/tok/document?preview=1",
                                           guest=True)
        self.assertNotIn("<img src=x", html)
        self.assertIn("&lt;img", html)
        self.assertIn("Shared link", html)

    def test_meta_line_reports_type_name_size_and_pages(self):
        html = reports._document_page_html(self._row(), file_url="u", preview_url="p")
        self.assertIn("PowerPoint", html)
        self.assertIn("Q3_Deck.pptx", html)
        self.assertIn("2.0 KB", html)
        self.assertIn("12 pages", html)

    # ── embed: the framed copy must not draw a second header ─────────────────
    #
    # The Workspace viewer loads this very page in an iframe and has already drawn a
    # title, a meta line and a download button above it. The standalone page is correct
    # to carry all three — it IS the interface when a colleague opens the link. So the two
    # shapes are pinned from both sides: adding the header back, or dropping it from the
    # standalone page, each fails a different assertion below.

    def test_the_standalone_page_carries_the_header(self):
        html = reports._document_page_html(self._row(), file_url="u", preview_url="p")
        self.assertIn('<header class="dp-bar">', html)
        self.assertIn("Q3 deck", html)          # the title
        self.assertIn("Download", html)         # the download
        self.assertIn("dp-meta", html)          # the meta line

    def test_the_framed_page_drops_the_whole_header(self):
        # ⚠️ Every part of it, not just the title: the shell above the frame shows the
        # filename, the state and the size, and it showed "37 KB" while this page said
        # "37.2 KB". Two headers is not redundancy, it is the reader having to work out
        # which number is true.
        #
        # ⚠️ Asserted on the BODY, not on the whole file. One template serves both shapes,
        # so `.dp-title` / `.dp-meta` / `.dp-dl` have to stay in the `<style>` block for
        # the standalone page — and a grep of the full HTML reports them as present on a
        # page that renders no header at all. Only markup can answer this.
        html = reports._document_page_html(self._row(), file_url="u", preview_url="p", embed=True)
        body = html.split("<body>", 1)[1]
        self.assertNotIn('<header class="dp-bar">', body)
        for gone in ("dp-title", "dp-meta", "dp-dl", "dp-spacer", "dp-note"):
            self.assertNotIn(gone, body, gone)
        # The document itself is still there — this is chrome removal, not a blank page.
        self.assertIn('<iframe class="dp-frame"', body)
        self.assertIn("Right-click", body)

    def test_the_framed_page_gives_its_full_height_to_the_document(self):
        # Without the header above it, reserving 96px for a bar that is not there would
        # leave a permanent empty strip under every embedded document.
        html = reports._document_page_html(self._row(), file_url="u", preview_url="p", embed=True)
        self.assertIn("height:100vh", html)
        self.assertNotIn("calc(100vh - 96px)", html)

    def test_the_frame_flag_does_not_reach_the_shared_link(self):
        # `embed` is a query flag on an address that has to stay the one stable
        # permalink a colleague can be sent. `_document_page_response` only forwards it;
        # nothing appends it to `download_url`, so the copy a colleague opens is always
        # the full page.
        html = reports._document_page_html(self._row(), file_url="s/tok/document",
                                           preview_url="p", embed=True)
        self.assertNotIn("embed", html)


class DocumentMetaTests(unittest.TestCase):
    def test_a_document_card_reports_its_type_name_and_pages(self):
        row = {"id": 1, "slug": "deck", "title": "Deck", "kind": "document", "doc_type": "pptx",
               "doc_name": "Deck.pptx", "doc_pages": 9, "has_filters": False,
               "owner_email": "a@example.com", "visibility": "private", "status": "published",
               "created_at": datetime(2026, 9, 30), "updated_at": datetime(2026, 9, 30)}
        meta = reports._meta(row, "a@example.com")
        self.assertEqual(meta["kind"], "document")
        self.assertEqual(meta["doc_type"], "pptx")
        self.assertEqual(meta["doc_name"], "Deck.pptx")
        self.assertEqual(meta["doc_pages"], 9)
        self.assertEqual(meta["langs"], [])
        self.assertTrue(meta["can_manage"])

    def test_a_dynamic_card_advertises_its_sidebar_without_shipping_the_schema(self):
        row = {"id": 2, "slug": "dyn", "title": "Dyn", "kind": "dynamic", "has_filters": True,
               "filters_json": '{"version":1,"filters":[{"key":"p"}]}',
               "owner_email": "a@example.com", "visibility": "private", "status": "published",
               "created_at": datetime(2026, 9, 30), "updated_at": datetime(2026, 9, 30)}
        meta = reports._meta(row, "a@example.com")
        self.assertEqual(meta["kind"], "dynamic")
        self.assertTrue(meta["has_filters"])
        self.assertNotIn("filters_json", meta)
        self.assertEqual(meta["doc_type"], "")


class DocumentPreviewRouteTests(unittest.TestCase):
    """The preview routes, wired: real bytes in, the reader's plan out.

    `doc_preview` is exercised on its own elsewhere; what these cover is the WIRING — which
    route a type points at, that a parse failure degrades to a download instead of a 500,
    and that the pptx/docx PDF cache is keyed so re-uploading cannot serve the old render.
    """

    OWNER = "zhen.huang@example.com"

    def setUp(self):
        self.row = {"id": 11, "slug": "deck", "title": "Deck", "status": "published",
                    "owner_email": self.OWNER, "visibility": "private", "doc_type": "pptx",
                    "doc_object": "deck-20260930-2200-abcd1234.pptx",
                    "doc_mime": reports.DOC_MIME["pptx"], "doc_name": "Deck.pptx",
                    "doc_pages": 2, "size_bytes": 1000,
                    "updated_at": datetime(2026, 9, 30, 22, tzinfo=timezone.utc)}
        self.cache_writes = []

    def _patches(self, data: bytes, row=None):
        row = row or self.row

        def _get(prefix, key):
            if key.startswith("preview/"):
                raise FileNotFoundError(key)                 # cache miss
            return data

        def _put(prefix, key, payload, content_type=None):
            self.cache_writes.append((key, len(payload), content_type))
            return prefix + "/" + key

        return [
            # `@_with_schema` creates the table on first use — that needs a database, and
            # nothing here is about the database.
            patch.object(reports, "_ensure_table", lambda: None),
            patch.object(reports, "_document_row", lambda slug, request=None: dict(row)),
            patch.object(reports, "_identity", lambda request: (self.OWNER, "browser")),
            patch.object(reports.oss_storage, "get_object", _get),
            patch.object(reports.oss_storage, "put_object", _put),
        ]

    def _preview(self, data, row=None, doc_type=None):
        row = dict(row or self.row)
        if doc_type:
            row["doc_type"] = doc_type
            row["doc_mime"] = reports.DOC_MIME[doc_type]
            row["doc_object"] = "x." + doc_type
        patches = self._patches(data, row)
        for patcher in patches:
            patcher.start()
        try:
            return asyncio.run(reports.get_document_preview("deck", None))
        finally:
            for patcher in patches:
                patcher.stop()

    def test_a_pdf_card_points_at_the_html_preview(self):
        # A PDF is planned as `html` now: the preview is where the note layer lives, and a
        # browser PDF plugin is a black box we cannot annotate (2026-10-01).
        plan = self._preview(_pdf_bytes(), doc_type="pdf")
        self.assertEqual(plan["kind"], "html")
        self.assertTrue(plan["url"].endswith("/api/reports/deck/document/preview.html"))

    def test_a_pptx_card_points_at_the_html_preview(self):
        plan = self._preview(_pptx_bytes(2), doc_type="pptx")
        self.assertEqual(plan["kind"], "html")
        self.assertTrue(plan["url"].endswith("/api/reports/deck/document/preview.html"))

    def test_a_workbook_is_planned_as_a_grid_and_carries_its_rows(self):
        plan = self._preview(_xlsx_bytes(), doc_type="xlsx", row={**self.row, "doc_type": "xlsx"})
        # `html` because the reader gets the Excel-shaped grid; `sheets` because a client may
        # want the rows themselves and should not have to parse the markup to get them.
        self.assertEqual(plan["kind"], "html")
        self.assertTrue(plan["url"].endswith("/api/reports/deck/document/preview.html"))
        # A hidden sheet and an empty one are both skipped: the reader is looking at the
        # workbook's content, and an empty grid tells them nothing.
        self.assertEqual([sheet["name"] for sheet in plan["sheets"]], ["Q3", "Q4"])
        self.assertEqual(plan["sheets"][0]["columns"], ["channel", "gto"])
        self.assertEqual(plan["sheets"][0]["rows"][0], ["online", 1200])

    def test_unrenderable_bytes_degrade_to_a_download_instead_of_a_500(self):
        plan = self._preview(b"PK\x03\x04 not really a deck", doc_type="pptx")
        self.assertEqual(plan["kind"], "download")
        self.assertIn("download=1", plan["url"])
        self.assertIn("could not be rendered", plan["reason"])

    def test_the_pptx_preview_is_rendered_once_and_cached_under_a_key_that_moves(self):
        try:
            import playwright  # noqa: F401
        except ImportError:                                   # pragma: no cover
            self.skipTest("playwright is not installed")
        data = _pptx_bytes(2)
        patches = self._patches(data)
        for patcher in patches:
            patcher.start()
        try:
            pdf = asyncio.run(reports.get_document_preview_pdf("deck", None))
        except HTTPException as exc:                          # pragma: no cover — no chromium
            self.skipTest("chromium unavailable: " + str(exc.detail))
        finally:
            for patcher in patches:
                patcher.stop()
        self.assertTrue(pdf.media_type == "application/pdf")
        self.assertTrue(pdf.body.startswith(b"%PDF"))
        self.assertEqual(len(self.cache_writes), 1)
        key, size, content_type = self.cache_writes[0]
        # The key carries the row's updated_at: re-uploading the file must not serve the
        # previous rendering, which is the whole reason it is in the key.
        self.assertIn(str(int(self.row["updated_at"].timestamp())), key)
        self.assertEqual(size, len(pdf.body))
        self.assertEqual(content_type, "application/pdf")


class _DocCursor:
    """Captures the publish INSERT and answers the follow-up read the response needs."""

    def __init__(self, existing=None, saved_slug="q3-deck"):
        self.existing = existing
        self.saved_slug = saved_slug
        self.insert = None
        self._row = None
        self.description = None

    def execute(self, sql, params=None):
        text = " ".join(str(sql).split()).lower()
        params = params or ()
        self._row = None
        if text.startswith("select") and "from ai_reports where slug" in text:
            self._row = self.existing
        elif text.startswith("insert into ai_reports"):
            self.insert = (text, params)
            self._row = {"id": 42, "created_at": datetime(2026, 9, 30, 22, tzinfo=timezone.utc),
                         "updated_at": datetime(2026, 9, 30, 22, tzinfo=timezone.utc),
                         "slug": params[0]}
        elif "order by pinned" in text or "distinct category" in text:
            self._row = None
        return self

    def fetchone(self):
        return self._row

    def fetchall(self):
        return [self._row] if self._row else []

    def __enter__(self):
        return self

    def __exit__(self, *exc):
        return False

    def close(self):
        pass


class PublishDocumentTests(unittest.TestCase):
    """The agent-facing upload endpoint itself: guards, the stored row, and the links."""

    OWNER = "zhen.huang@example.com"

    def _publish(self, body, existing=None, saved_slug="q3-deck"):
        cursor = _DocCursor(existing=existing, saved_slug=saved_slug)
        conn = type("C", (), {"cursor": lambda self, **k: cursor, "commit": lambda self: None,
                              "close": lambda self: None, "__enter__": lambda self: self,
                              "__exit__": lambda self, *e: False})()
        uploads, removals = [], []
        patchers = [
            patch.object(reports, "_ensure_table", lambda: None),
            patch.object(reports, "_pg", lambda: conn),
            patch.object(reports, "_identity", lambda request: (self.OWNER, "agent")),
            patch.object(reports.oss_storage, "put_object",
                         lambda prefix, key, data, content_type=None: uploads.append((key, len(data))) or key),
            patch.object(reports.oss_storage, "remove_object",
                         lambda prefix, key: removals.append(key)),
            patch.object(reports, "_cover_from_body", lambda body, title, request=None: None),
            patch.object(reports, "_store_cover", lambda cur, slug, cover, clear=False: None),
            patch.object(reports, "_public_report_url", lambda request, slug: "/r/" + slug),
            patch.object(reports, "_absolute_url", lambda request, path: path),
            # The response reads the row back through `_load`; feed it the row the INSERT
            # would have produced, so `_meta` shapes a real payload.
            patch.object(reports, "_load", lambda slug, include_html=False, viewer_email="",
                         request=None:
                         dict(_DOC_META_ROW, slug=slug)),
        ]
        for patcher in patchers:
            patcher.start()
        try:
            return asyncio.run(reports.publish_document(body, None)), cursor, uploads, removals
        finally:
            for patcher in patchers:
                patcher.stop()

    def test_a_deck_becomes_a_document_card_with_both_links(self):
        raw = _pptx_bytes(2)
        body = reports.DocumentIn(filename="Q3_Deck.pptx", title="Q3 deck", slug="q3-deck",
                                 content_base64=base64.b64encode(raw).decode("ascii"),
                                 submitter="Zhen Huang")
        meta, cursor, uploads, removals = self._publish(body)
        text, params = cursor.insert
        self.assertIn("'document'", text)
        self.assertEqual(params[0], "q3-deck")                       # the caller's slug, verbatim
        # (`html` and `langs` are literals in the statement — a document has neither.)
        self.assertEqual(params[12], len(raw))                       # size_bytes
        self.assertEqual(params[13], self.OWNER)                     # owner comes from the credential
        self.assertEqual(params[14], "pptx")
        self.assertTrue(params[15].endswith(".pptx"))                # the object key keeps the type
        self.assertEqual(params[17], "Q3_Deck.pptx")                 # doc_name
        self.assertEqual(params[18], 2)                              # doc_pages, from the real file
        self.assertEqual(uploads, [(params[15], len(raw))])           # uploaded exactly once
        self.assertEqual(removals, [])
        self.assertEqual(meta["kind"], "document")
        self.assertEqual(meta["url"], "/r/q3-deck")
        # `_absolute_url` is stubbed in this harness, so the mount point is what the real
        # `_base_href()` contributes — the path shape is what matters here.
        self.assertEqual(meta["document_url"], "/api/reports/q3-deck/document")
        self.assertEqual(meta["download_url"], "/api/reports/q3-deck/document?download=1")
        self.assertEqual(meta["format"], "document")

    def test_the_two_slugs_that_belong_to_someone_else_are_refused_before_any_upload(self):
        # An HTML report on that slug.
        body = reports.DocumentIn(filename="deck.pptx", slug="q3-report",
                                 content_base64=base64.b64encode(_pptx_bytes(1)).decode("ascii"))
        with self.assertRaises(HTTPException) as ctx:
            self._publish(body, existing={"id": 1, "doc_object": "", "doc_type": "",
                                          "owner_email": self.OWNER, "visibility": "private"})
        self.assertEqual(ctx.exception.status_code, 409)
        # A card owned by somebody else (or a public snapshot of one).
        with self.assertRaises(HTTPException) as ctx:
            self._publish(body, existing={"id": 2, "doc_object": "x.pptx", "doc_type": "pptx",
                                          "owner_email": "other@example.com", "visibility": "private"})
        self.assertEqual(ctx.exception.status_code, 409)

    def test_a_legacy_extension_never_reaches_storage(self):
        body = reports.DocumentIn(filename="old.ppt", content_base64=base64.b64encode(b"x").decode("ascii"))
        with self.assertRaises(HTTPException) as ctx:
            self._publish(body)
        self.assertEqual(ctx.exception.status_code, 400)
        self.assertIn(".pptx", str(ctx.exception.detail))

    def test_a_missing_slug_always_allocates_a_new_suffixed_address(self):
        body = reports.DocumentIn(content_base64=base64.b64encode(_pptx_bytes(1)).decode("ascii"),
                                 filename="deck.pptx")
        _meta, cursor, _uploads, _removals = self._publish(body)
        slug = cursor.insert[1][0]
        self.assertTrue(slug.startswith("deck-"), slug)
        self.assertRegex(slug, r"-\d{8}-\d{4}-[a-z0-9]{6}$")


_DOC_META_ROW = {
    "id": 42, "slug": "q3-deck", "title": "Q3 deck", "kind": "document", "doc_type": "pptx",
    "doc_name": "Q3_Deck.pptx", "doc_pages": 2, "has_filters": False, "status": "published",
    "owner_email": PublishDocumentTests.OWNER, "visibility": "private",
    "created_at": datetime(2026, 9, 30, 22, tzinfo=timezone.utc),
    "updated_at": datetime(2026, 9, 30, 22, tzinfo=timezone.utc),
}


if __name__ == "__main__":
    unittest.main()

class DocumentReaderMetadataTests(unittest.TestCase):
    """The reader page's meta line must describe the file it is showing.

    ⚠️ Regression guard, 2026-10-01. Both document-page data sources — `render_report`
    (`/r/{slug}`) and `_resolve_anyone_link` (`/s/{token}`) — SELECT a fixed column list
    instead of `*`, and BOTH omitted `size_bytes` and `doc_pages`. Nothing failed: the row
    was fine, `row.get(...)` was `None`, `_human_size(None)` returns "0 B", and so every
    document's reader page quietly advertised a zero-byte file with no page count. It was
    found by uploading one of each type and reading the page, not by a test — hence this
    one, which reads the two queries instead of hoping someone notices the next time a
    column is added to the template.
    """

    def test_both_document_page_sources_select_the_meta_fields(self):
        import inspect

        from routers import reports

        # ⚠️ The list is what the reader page actually reads off the row. `slug` was the one
        # that broke everything: the page builds the document's own addresses from it, so an
        # unselected slug produced `api/reports//document`, which 404'd every preview and made
        # every download save the error JSON as `document.json` (2026-10-01).
        for fn in (reports.render_report, reports._resolve_anyone_link):
            source = inspect.getsource(fn)
            for column in ("slug", "size_bytes", "doc_pages", "doc_type", "doc_name", "title"):
                self.assertIn(column, source,
                              "%s must select %s — the document reader page uses it"
                              % (fn.__name__, column))

    def test_the_document_urls_are_built_from_the_row_slug(self):
        # The page's own addresses must not contain an empty slug — that is the shape of the
        # 2026-10-01 breakage, and it is cheap to assert directly.
        from routers import reports

        for row in ({"slug": "deck", "title": "T", "doc_type": "pdf", "doc_name": "d.pdf",
                     "size_bytes": 10, "doc_pages": 1},
                    {"slug": "sheet", "title": "T", "doc_type": "xlsx", "doc_name": "s.xlsx",
                     "size_bytes": 10, "doc_pages": 1}):
            page = reports._document_page_html(
                row, file_url="api/reports/%s/document" % row["slug"],
                preview_url="api/reports/%s/document/preview.pdf" % row["slug"])
            self.assertNotIn("//document", page)
            self.assertIn('api/reports/%s/document' % row["slug"], page)

    def test_the_meta_line_reports_a_real_size(self):
        from routers.reports import _document_page_html

        page = _document_page_html(
            {"slug": "d", "title": "Deck", "doc_type": "pptx", "doc_name": "deck.pptx",
             "size_bytes": 34771, "doc_pages": 7},
            file_url="api/reports/d/document", preview_url="api/reports/d/document/preview.pdf")
        self.assertIn("34.0 KB", page)
        self.assertIn("7 pages", page)
        self.assertNotIn("· 0 B", page)

    def test_a_single_page_or_sheet_is_not_plural(self):
        # A real one-sheet workbook read "1 sheets" before 2026-10-01.
        from routers.reports import _document_page_html

        sheet = _document_page_html(
            {"slug": "d", "title": "Book", "doc_type": "xlsx", "doc_name": "b.xlsx",
             "size_bytes": 5357, "doc_pages": 1},
            file_url="api/reports/d/document", preview_url="api/reports/d/document/preview.pdf")
        self.assertIn("1 sheet", sheet)
        self.assertNotIn("1 sheets", sheet)
        many = _document_page_html(
            {"slug": "d", "title": "Book", "doc_type": "xlsx", "doc_name": "b.xlsx",
             "size_bytes": 5357, "doc_pages": 3},
            file_url="api/reports/d/document", preview_url="api/reports/d/document/preview.pdf")
        self.assertIn("3 sheets", many)


class PreviewFontTests(unittest.TestCase):
    """Which preview carries a font, and which one deliberately does not.

    ⚠️ 2026-10-01: a pptx/docx preview used to be a server-rendered PDF, and the API image ships
    no CJK font at all (`api/fonts/` holds a README), so every Chinese character came out as a
    box. The fix has two halves and they must stay on opposite sides: the markup served to a
    BROWSER stays lean (the reader's machine has fonts), and the server-side PDF render inlines
    one (that machine does not).
    """

    @staticmethod
    def _docx_with_chinese() -> bytes:
        import io as _io

        from docx import Document

        document = Document()
        document.add_heading("渠道进销存口径表", level=1)
        document.add_paragraph("甲公司、乙公司、丙公司、丁公司、戊公司")
        buffer = _io.BytesIO()
        document.save(buffer)
        return buffer.getvalue()

    def test_the_served_markup_has_no_inlined_font(self):
        from services import doc_preview

        html = doc_preview.preview_html(self._docx_with_chinese(), "docx", "T")
        self.assertIn("渠道", html)                 # the text is there …
        self.assertNotIn("@font-face", html)        # … and no 1.4MB of base64 rides along

    def test_the_pdf_render_is_the_side_that_inlines_it(self):
        import inspect

        from services import doc_preview

        source = inspect.getsource(doc_preview.html_to_pdf)
        self.assertIn("_cjk_font_face()", source,
                      "html_to_pdf must inline the CJK face: it renders inside an image with none")
        face = doc_preview._cjk_font_face()
        self.assertIn("data:font/woff2;base64,", face)
        self.assertIn("@font-face", face)


def _docx_with_page_break() -> bytes:
    """Two sentences, one explicit ``w:br type=page`` between them."""
    import io as _io

    from docx import Document
    from docx.enum.text import WD_BREAK

    document = Document()
    document.add_paragraph("first page")
    paragraph = document.add_paragraph()
    paragraph.add_run().add_break(WD_BREAK.PAGE)
    document.add_paragraph("second page")
    buffer = _io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


class OfficeLayoutTests(unittest.TestCase):
    """A preview has to look like the application the file came from.

    Asked for 2026-10-01: "docx和pptx，以及xlsx文件也要还原pages，slide和sheet的视觉效果，
    让用户看起来像是在office里看一样". These assert the load-bearing parts of that — the ones a
    later CSS tidy-up would quietly delete: the grey desk, the white sheets, the page/slide
    split, Excel's column letters and row numbers.
    """

    def test_a_deck_is_a_stack_of_slide_cards(self):
        from services import doc_preview

        html = doc_preview.preview_html(_pptx_bytes(3), "pptx", "T", max_slides=10)
        self.assertIn('class="pv-canvas pv-canvas--slides"', html)
        self.assertEqual(html.count('class="pv-slide pv-sheet-frame"'), 3)
        self.assertEqual(html.count('class="pv-slide-wrap"'), 3)
        self.assertIn("Slide 1", html)
        self.assertIn("Slide 3", html)

    def test_a_docx_is_sheets_of_paper_split_on_the_files_own_page_breaks(self):
        from services import doc_preview

        html = doc_preview.preview_html(_docx_with_page_break(), "docx", "T")
        self.assertEqual(html.count('class="pv-page pv-sheet-frame"'), 2)
        self.assertIn("Page 1", html)
        self.assertIn("Page 2", html)
        # The break marker itself must not survive into the markup: it is a split point, not
        # an element, and leaving one behind would print a stray break inside the sheet.
        # (Asserted on the BODY: `pv-pagebreak` also appears in the stylesheet, where it is a
        #  belt-and-braces `display:none` for a marker that ever escapes the split.)
        body = html.split("</head>", 1)[1]
        self.assertNotIn('class="pv-pagebreak"', body)
        first, second = html.split("Page 2")
        self.assertIn("first page", first)
        self.assertIn("second page", second)

    def test_a_docx_without_page_breaks_is_one_sheet(self):
        from services import doc_preview

        html = doc_preview.preview_html(PreviewFontTests._docx_with_chinese(), "docx", "T")
        self.assertEqual(html.count('class="pv-page pv-sheet-frame"'), 1)

    def test_a_trailing_page_break_does_not_add_a_blank_sheet(self):
        # Word does render the page a break pushes you onto; a break at the very END of the
        # file is not that, and rendering it would put a permanently blank sheet in the scroll.
        import io as _io

        from docx import Document
        from docx.enum.text import WD_BREAK

        document = Document()
        document.add_paragraph("only page")
        document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        buffer = _io.BytesIO()
        document.save(buffer)

        from services import doc_preview

        html = doc_preview.preview_html(buffer.getvalue(), "docx", "T")
        self.assertEqual(html.count('class="pv-page pv-sheet-frame"'), 1)
        self.assertIn("only page", html)

    def test_a_workbook_is_a_grid_with_excels_own_chrome(self):
        from services import doc_preview

        html = doc_preview.tables_html(_xlsx_bytes(), "T")
        self.assertIn('class="pv-cells"', html)
        # Column letters, not column headings: A and B above the two columns of the fixture.
        self.assertIn('<th class="pv-colhead">A</th>', html)
        self.assertIn('<th class="pv-colhead">B</th>', html)
        # Row numbers: the fixture's first row IS the header, so it is Excel row 1 and the
        # first data row is 2. Numbering the data from 1 would misname every row a reader is
        # ever asked to look at.
        self.assertIn('<th class="pv-gutter">1</th>', html)
        self.assertIn('<th class="pv-gutter">2</th>', html)
        self.assertNotIn('<th class="pv-gutter">0</th>', html)
        # The sheet names ride in Excel's tab strip, and the hidden/empty sheets are absent.
        self.assertIn('class="pv-tabs"', html)
        self.assertIn('class="pv-tab is-on"', html)
        self.assertNotIn("Hidden", html)
        self.assertNotIn("Empty", html)

    def test_a_workbook_keeps_its_thousands_separators_and_its_years(self):
        from services import doc_preview

        import io as _io

        from openpyxl import Workbook

        book = Workbook()
        sheet = book.active
        sheet.append(["year", "gto", "small"])
        sheet.append([2026, 18420000, 5000])
        sheet["C2"].number_format = "#,##0"       # the file asking for grouping, out loud
        buffer = _io.BytesIO()
        book.save(buffer)
        html = doc_preview.tables_html(buffer.getvalue(), "T")
        self.assertIn(">18,420,000<", html)       # plain number, big enough to need grouping
        self.assertIn(">5,000<", html)            # cell format wins even under the size floor
        self.assertIn(">2026<", html)             # a year is a different fact, never "2,026"
        self.assertNotIn(">2,026<", html)

    def test_a_pdf_becomes_pages_of_paper_with_its_pages_as_images(self):
        from services import doc_preview

        html = doc_preview.preview_html(_pdf_bytes(3), "pdf", "T", max_slides=10)
        self.assertEqual(html.count('class="pv-page pv-page--pdf pv-sheet-frame"'), 3)
        self.assertEqual(html.count("data:image/png;base64,"), 3)
        self.assertIn("Page 1", html)
        self.assertIn("Page 3", html)

    def test_a_pdf_stops_at_the_page_cap_and_says_so(self):
        from services import doc_preview

        # `max_pages` is a parameter, not the module constant: the constant is only the default
        # the routes use, and patching it here would not reach the already-bound argument.
        html = doc_preview._pdf_html(_pdf_bytes(5), "T", max_pages=2)
        self.assertEqual(html.count("data:image/png;base64,"), 2)
        self.assertIn("showing the first 2 of 5 pages", html)

    def test_every_preview_names_the_pages_a_note_can_land_on(self):
        # `PAGE_SELECTOR` is the contract between the renderer and the note runtime: it has to
        # match the markup above, or notes silently fall back to "one long page".
        from services import doc_preview

        selector = doc_preview.PAGE_SELECTOR
        for kind, data in (("pptx", _pptx_bytes(2)), ("docx", _docx_with_page_break()),
                           ("xlsx", _xlsx_bytes()), ("pdf", _pdf_bytes(2))):
            html = doc_preview.preview_html(data, kind, "T")
            for token in selector.split(", "):
                if 'class="' + token[1:] in html:
                    break
            else:
                self.fail(f"{kind}: no element matches PAGE_SELECTOR {selector!r}")


class PreviewNoteLayerTests(unittest.TestCase):
    """A document card's preview has to arrive with its note layer attached.

    Asked for 2026-10-01: "这些类型的文档也需要可以加notes，与之前html和md文档的要求一样". The
    renderer knows nothing about notes — this is the wiring that adds them, and it is the part
    a refactor would silently drop (the file would still preview perfectly, just unannotatable).
    """

    ROW = {"slug": "q3-deck", "title": "Q3 deck", "doc_type": "pptx", "kind": "document"}

    def test_the_layer_is_inlined_with_the_page_selector(self):
        html = reports._preview_with_notes("<html><head></head><body><p>x</p></body></html>",
                                           self.ROW, pages=".pv-page, .pv-slide, .pv-sheet",
                                           viewer_email="reader@example.com")
        self.assertIn("data-doc-annotations", html)
        self.assertIn("window.DocAnnotations", html)
        # The page selector is what makes a note belong to slide 4 instead of "62% down a scroll".
        self.assertIn('"pages": ".pv-page, .pv-slide, .pv-sheet"', html)
        self.assertIn('"writable": true', html)
        self.assertIn("reader@example.com", html)

    def test_the_layer_draws_its_icons_from_remixicon(self):
        # The preview document has no theme of its own; without this the note toggle and its
        # right-click menu render as empty buttons.
        html = reports._preview_with_notes("<html><head></head><body></body></html>",
                                           self.ROW, pages=".pv-slide")
        self.assertIn("remixicon/remixicon.css", html)
        self.assertLess(html.index("remixicon"), html.index("data-doc-annotations"))

    def test_a_guest_link_gets_the_layer_but_no_write_path(self):
        # A note is signed by a session and shows up in a colleague's Inbox; a share token is
        # not a session. A guest may read the notes that are already there and never add one.
        html = reports._preview_with_notes("<html><head></head><body></body></html>",
                                           self.ROW, pages=".pv-slide", writable=False,
                                           read_url="/s/tok/annotations")
        self.assertIn('"writable": false', html)
        self.assertIn("/s/tok/annotations", html)

    def test_an_empty_preview_is_left_alone(self):
        self.assertEqual(reports._preview_with_notes("", self.ROW, pages=".pv-slide"), "")


class RenderReportEmbedRouteTests(unittest.TestCase):
    """`GET /r/{slug}?embed=1`, wired: the query flag has to REACH the template.

    ⚠️ This class exists because of a mutation, not because of a bug. `DocumentPageTests`
    calls `_document_page_html(..., embed=True)` directly, so all of those assertions kept
    passing after `render_report` was changed to drop the flag on the floor — the template
    understood `embed` and the route never mentioned it. That is the whole feature: the
    iframe's address carries `?embed=1`, and a template that is never told is a template
    that still draws the second header. So the flag is asserted where it is READ.
    """

    OWNER = "zhen.huang@example.com"
    ROW = {"id": 11, "slug": "deck", "html": None, "title": "Deck", "status": "published",
           "owner_email": OWNER, "visibility": "private", "kind": "document",
           "filters_json": None, "doc_type": "pptx", "doc_object": "deck.pptx",
           "doc_mime": "application/vnd.ms-powerpoint", "doc_name": "Deck.pptx",
           "doc_pages": 2, "size_bytes": 1000}

    class _Cursor:
        def __init__(self, row):
            self._row = row

        def __enter__(self):
            return self

        def __exit__(self, *exc):
            return False

        def execute(self, sql, params=None):
            self._row["_last_sql"] = sql

        def fetchone(self):
            return dict(self._row)

    class _Conn:
        def __init__(self, row):
            self._row = row

        def __enter__(self):
            return self

        def __exit__(self, *exc):
            return False

        def cursor(self, cursor_factory=None):
            return RenderReportEmbedRouteTests._Cursor(self._row)

    def _render(self, **kwargs):
        row = dict(self.ROW)
        # ⚠️ `download` is passed EXPLICITLY. Called directly, the route's signature
        # defaults are `Query(False)` objects, not booleans — and a `Query` instance is
        # truthy, so the omitted form takes the download branch and every "page" assertion
        # below would be reading the body of a 302.
        kwargs.setdefault("download", False)
        patches = [
            patch.object(reports, "_ensure_table", lambda: None),
            patch.object(reports, "_db", lambda: self._Conn(row)),
            patch.object(reports, "_identity", lambda request: (self.OWNER, "browser")),
            patch.object(reports, "_may_read", lambda row, viewer: True),
            patch.object(reports, "_inject_base_tag", lambda html, href: html),
            patch.object(reports, "_base_href", lambda: "/"),
        ]
        for patcher in patches:
            patcher.start()
        try:
            response = asyncio.run(reports.render_report("deck", None, **kwargs))
            return response.body.decode("utf-8")
        finally:
            for patcher in patches:
                patcher.stop()

    def test_the_query_flag_reaches_the_template(self):
        body = self._render(embed=True)
        self.assertNotIn('<header class="dp-bar">', body)
        self.assertIn('<iframe class="dp-frame"', body)

    def test_without_the_flag_the_page_is_unchanged(self):
        # The default is the shape a colleague gets by opening the link, and it is the
        # shape that has to keep working — `embed` is an addition to the page, not a
        # replacement of it.
        body = self._render(embed=False)
        self.assertIn('<header class="dp-bar">', body)
        self.assertIn("Download", body)
