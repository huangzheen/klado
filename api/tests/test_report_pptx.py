import io
import struct
import unittest
from zipfile import ZipFile

from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN

from services.report_pptx import UnsupportedReport, build_pptx


def rect(kind, x=100, y=100, w=300, h=100, **extra):
    return {"kind": kind, "x": x, "y": y, "w": w, "h": h, **extra}


class ReportPptxTests(unittest.TestCase):
    def test_bundled_faces_travel_with_editable_text(self):
        output = build_pptx({"slides": [{"nodes": [
            rect("text", text="Hello 中文", style={"font": "Roboto Condensed, CoolSans SC Narrow", "size": 20}),
        ]}]})
        with ZipFile(io.BytesIO(output)) as archive:
            font_parts = sorted(n for n in archive.namelist() if n.startswith("ppt/fonts/"))
            self.assertEqual(len(font_parts), 4)  # regular and bold in each family
            presentation = archive.read("ppt/presentation.xml")
            self.assertIn(b'embedTrueTypeFonts="1"', presentation)
            self.assertIn(b'saveSubsetFonts="0"', presentation)
            self.assertIn(b'CoolSans SC Narrow', presentation)
            self.assertIn(b'Roboto Condensed', presentation)
            for part in font_parts:
                payload = archive.read(part)
                self.assertEqual(struct.unpack_from("<I", payload)[0], len(payload))
                self.assertEqual(struct.unpack_from("<H", payload, 34)[0], 0x504C)
        self.assertEqual(Presentation(io.BytesIO(output)).slides[0].shapes[0].text, "Hello 中文")

    def test_writes_editable_text_shape_table_and_chart(self):
        cell = {"x": 100, "y": 300, "w": 200, "h": 50, "text": "Revenue",
                "rowspan": 1, "colspan": 1,
                "style": {"font": "Arial", "size": 20, "color": "#123456",
                          "background": "#ffffff", "align": "left"}}
        scene = {"slides": [{"background": "#ffffff", "nodes": [
            rect("shape", fill="#2563eb", radius=8),
            rect("text", y=210, text="Editable title",
                 style={"font": "Arial", "size": 40, "color": "#123456", "bold": True}),
            rect("table", y=300, w=200, h=50, rows=[[cell]]),
            rect("chart", y=400, w=500, h=300,
                 data={"type": "bar", "categories": ["Jan", "Feb"],
                       "series": [{"name": "Sales", "values": [10, 20], "color": "#2563eb"}]}),
        ]}]}
        output = Presentation(io.BytesIO(build_pptx(scene)))
        self.assertEqual(len(output.slides), 1)
        shapes = output.slides[0].shapes
        self.assertTrue(any(s.has_text_frame and "Editable title" in s.text for s in shapes))
        self.assertTrue(any(s.has_table and s.table.cell(0, 0).text == "Revenue" for s in shapes))
        self.assertTrue(any(s.has_chart for s in shapes))
        self.assertFalse(any(s.shape_type == 13 for s in shapes))  # no slide screenshot

    def test_rejects_documents_without_deck_pages(self):
        with self.assertRaisesRegex(UnsupportedReport, "Report Decks"):
            build_pptx({"slides": []})

    def test_mixed_style_list_stays_in_one_editable_textbox(self):
        normal = {"font": "Arial", "size": 20, "color": "#506070", "lineHeight": 34}
        bold = {**normal, "font": "Georgia", "bold": True, "color": "#0d1b2a"}
        scene = {"slides": [{"nodes": [rect("text", w=800, h=120, wrap=True, paragraphs=[
            {"style": normal, "runs": [
                {"text": "Down across ", "style": normal},
                {"text": "nine months", "style": bold},
                {"text": ", not one quarter.", "style": normal},
            ]},
            {"style": normal, "runs": [
                {"text": "ASP fell ", "style": normal},
                {"text": "14.5%", "style": bold},
            ]},
        ])]}]}
        output = Presentation(io.BytesIO(build_pptx(scene)))
        textboxes = [s for s in output.slides[0].shapes if s.has_text_frame and s.text]
        self.assertEqual(len(textboxes), 1)
        paragraphs = textboxes[0].text_frame.paragraphs
        self.assertEqual(len(paragraphs), 2)
        self.assertEqual("".join(r.text for r in paragraphs[0].runs), "Down across nine months, not one quarter.")
        self.assertEqual([r.font.bold for r in paragraphs[0].runs], [False, True, False])
        self.assertEqual(paragraphs[0].runs[1].font.name, "Georgia")

    def test_merged_heatmap_table_stays_native_and_keeps_cell_fills(self):
        cell = lambda text, w, h, bg, **span: {
            "w": w, "h": h, "text": text, "style": {"background": bg},
            "rowspan": span.get("rowspan", 1), "colspan": span.get("colspan", 1),
        }
        scene = {"slides": [{"nodes": [rect("table", w=300, h=100, rowHeights=[50, 50], rows=[
            [cell("Region", 100, 100, "rgb(255, 255, 255)", rowspan=2),
             cell("Quarter", 200, 50, "rgb(234, 242, 252)", colspan=2)],
            [cell("Q1", 100, 50, "rgb(218, 232, 249)"),
             cell("Q2", 100, 50, "rgb(29, 111, 214)")],
        ])]}]}
        output = Presentation(io.BytesIO(build_pptx(scene)))
        table = next(shape.table for shape in output.slides[0].shapes if shape.has_table)
        self.assertEqual((len(table.rows), len(table.columns)), (2, 3))
        self.assertTrue(table.cell(0, 0).is_merge_origin)
        self.assertTrue(table.cell(0, 1).is_merge_origin)
        self.assertEqual(table.cell(1, 1).fill.fore_color.rgb, (218, 232, 249))
        self.assertEqual(table.cell(1, 2).fill.fore_color.rgb, (29, 111, 214))
        self.assertFalse(any(s.shape_type == 13 for s in output.slides[0].shapes))

    def test_annotation_outlines_remain_native_without_decorative_shadow(self):
        scene = {"slides": [{"nodes": [
            rect("shape", fill="#ffffff", shadow={"color": "rgba(20, 24, 31, 0.06)",
                                                   "x": 0, "y": 1, "blur": 3, "spread": 0}),
            rect("line", x=100, y=100, w=300, h=0, color="#dc2626", width=5),
        ]}]}
        output = Presentation(io.BytesIO(build_pptx(scene)))
        shapes = output.slides[0].shapes
        self.assertIsNone(shapes[0]._element.spPr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}effectLst"))
        self.assertEqual(shapes[1].line.color.rgb, (220, 38, 38))
        self.assertFalse(any(shape._element.xpath("./p:style") for shape in shapes))
        self.assertFalse(any(s.shape_type == 13 for s in shapes))

    def test_table_uses_html_cell_edges_without_office_theme(self):
        cell = {"x": 0, "y": 0, "w": 100, "h": 40, "text": "A", "rowspan": 1, "colspan": 1,
                "style": {"background": "rgb(242, 247, 251)", "size": 18},
                "padding": {"left": 7, "right": 5, "top": 3, "bottom": 2},
                "verticalAlign": "top",
                "borders": {"bottom": {"color": "rgb(229, 235, 243)", "width": 1}}}
        output = Presentation(io.BytesIO(build_pptx({"slides": [{"nodes": [
            rect("table", x=0, y=0, w=100, h=40, rows=[[cell]], rowHeights=[40]),
        ]}]})))
        shape = output.slides[0].shapes[0]
        self.assertFalse(shape._element.xpath(".//a:tableStyleId"))
        tc = shape.table.cell(0, 0)._tc
        self.assertIn('val="E5EBF3"', tc.get_or_add_tcPr().xml)
        self.assertIn("<a:noFill", tc.get_or_add_tcPr().xml)
        self.assertIn('w="6350"', tc.get_or_add_tcPr().xml)
        exported_cell = shape.table.cell(0, 0)
        self.assertEqual(exported_cell.margin_left, 44450)
        self.assertEqual(exported_cell.margin_top, 19050)
        self.assertEqual(exported_cell.vertical_anchor, MSO_ANCHOR.TOP)

    def test_badge_is_one_editable_shape_with_centered_text_and_thin_edge(self):
        output = Presentation(io.BytesIO(build_pptx({"slides": [{"nodes": [
            rect("badge", x=10, y=20, w=80, h=28, fill="#e7efff", radius=10,
                 stroke={"color": "#bcd4fe", "width": 1}, text="Badge", fit=True,
                 valign="middle", margins={"left": 8, "right": 8, "top": 2, "bottom": 2},
                 style={"font": "Arial", "size": 14, "align": "center", "color": "#1d4ed8"}),
        ]}]})))
        self.assertEqual(len(output.slides[0].shapes), 1)
        badge = output.slides[0].shapes[0]
        self.assertEqual(badge.text, "Badge")
        self.assertEqual(badge.text_frame.vertical_anchor, MSO_ANCHOR.MIDDLE)
        self.assertEqual(badge.text_frame.auto_size, MSO_AUTO_SIZE.NONE)
        self.assertEqual(badge.text_frame.margin_left, 0)
        self.assertEqual(badge.text_frame.margin_right, 0)
        self.assertNotIn("normAutofit", badge.text_frame._txBody.bodyPr.xml)
        self.assertEqual(badge.text_frame.paragraphs[0].alignment, PP_ALIGN.CENTER)
        self.assertEqual(badge.line.color.rgb, (188, 212, 254))
        self.assertEqual(badge.line.width, 6350)  # one CSS pixel = half a point

    def test_html_break_in_badge_is_a_real_editable_line_break(self):
        scene = {"slides": [{"nodes": [rect("badge", w=24, h=38,
            fill="#b08968", fit=True, valign="middle", paragraphs=[{
                "style": {"font": "CoolSans SC Narrow", "size": 14, "align": "center"},
                "runs": [{"text": "线", "style": {"size": 14}},
                         {"text": "\v", "style": {"size": 14}},
                         {"text": "下", "style": {"size": 14}}],
            }])]}]}
        output = build_pptx(scene)
        badge = Presentation(io.BytesIO(output)).slides[0].shapes[0]
        self.assertEqual(badge.text, "线\v下")
        self.assertEqual(badge.text_frame.paragraphs[0].alignment, PP_ALIGN.CENTER)
        self.assertIn(b"<a:br", badge._element.xml.encode())
        self.assertNotIn("_x000B_", badge._element.xml)


if __name__ == "__main__":
    unittest.main()
