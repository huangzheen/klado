"""
The four Office previews and the notes a reader leaves on them, in a real browser.

What this proves that the unit tests cannot: that the markup, the note runtime and the
browser's own layout agree with each other. Every claim below is measured off the RENDERED
page (a computed background colour, a bounding box, the page index a click produced), because
"looks like Office" is a statement about pixels and the server never sees them.

The harness is the same shape as `verify_dynamic_report_ui.py`: a real uvicorn app with the
real router, real injection order and a stub database, plus a browser that drives it.

    api/.venv312/bin/python api/tests/verify_office_notes_ui.py      # checks
    api/.venv312/bin/python api/tests/verify_office_notes_ui.py --serve   # leave it up to look at

⚠️ Three traps this file has already fallen into, do not re-introduce:

1. **`FRONTEND_DIR` must point at the repo's `frontend/out`.** There used to be a second copy
   under `api/` that a container build shipped; it is gone, and this file hard-coded that path
   for a while — it then failed with `Directory ... does not exist` before a single check ran.
   (The `[::1]` bind below is a leftover from the same era; it is harmless but no longer
   required, because nothing in the SPA inspects the hostname any more.)
2. **`render_report` must be called with an explicit `download=False`.** Called directly from
   Python its `Query(False)` default is a Query OBJECT — truthy — so the response carries
   `Content-Disposition: attachment` and Chromium turns the iframe navigation into a download.
3. **A `--serve` demo left running holds the port.** The checks then talk to THAT process while
   their counters live in this one, and fail for reasons that have nothing to do with the code.
"""
import io
import json
import os
import sys
import threading
import time
from datetime import datetime, timezone
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import uvicorn  # noqa: E402
from fastapi import FastAPI, Query, Request  # noqa: E402
from fastapi.responses import FileResponse, JSONResponse  # noqa: E402
from fastapi.staticfiles import StaticFiles  # noqa: E402
from playwright.sync_api import sync_playwright  # noqa: E402
from starlette.responses import Response  # noqa: E402

from routers import annotations as annotations_router  # noqa: E402
from routers import reports  # noqa: E402
from services import annotations, inbox, oss_storage  # noqa: E402

PORT = int(os.environ.get("OFFICE_UI_PORT", "18897"))
SERVE_PORT = int(os.environ.get("OFFICE_UI_SERVE_PORT", "18898"))
OWNER = "reader@example.com"           # the browser session that leaves the note
DOC_OWNER = "author@example.com"       # whose document it is, so the Inbox notification has a recipient
TOKEN = "aaaaaaaaaaaaaaaaaaaa.bbbbbbbbbbbbbbbbbbbbbb"
FRONTEND = Path(os.environ.get(
    "FRONTEND_DIR", str(Path(__file__).resolve().parents[2] / "frontend" / "out")))

FAILURES: list[str] = []


# ── the four files ───────────────────────────────────────────────────────────

def _pptx_bytes() -> bytes:
    from pptx import Presentation
    presentation = Presentation()
    for index in range(3):
        slide = presentation.slides.add_slide(presentation.slide_layouts[5])
        slide.shapes.title.text = f"Slide {index + 1} 渠道复盘"
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def _docx_bytes() -> bytes:
    from docx import Document
    from docx.enum.text import WD_BREAK
    document = Document()
    document.add_heading("产品营销月度例会 · 纪要", level=1)
    document.add_paragraph("时间 2026-10-01 10:00-11:30 · 地点 线上")
    paragraph = document.add_paragraph()
    paragraph.add_run().add_break(WD_BREAK.PAGE)
    document.add_heading("第二页 · 行动项", level=2)
    document.add_paragraph("补齐门店样机 / Wang Fang / 2026-11-14")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _xlsx_bytes() -> bytes:
    from openpyxl import Workbook
    book = Workbook()
    sheet = book.active
    sheet.title = "渠道口径"
    sheet.append(["渠道", "2026-08 销售额", "2026-09 销售额"])
    for row in (("Alpha Stores", 18420000, 19100000), ("Acme Retail", 12980000, 12150000),
                ("Beta Direct", 9640000, 10120000)):
        sheet.append(list(row))
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def _pdf_bytes() -> bytes:
    from services.pdf_io import text_pdf
    return text_pdf([[(72, 120, 18, f"One Pager page {index + 1}")] for index in range(2)])


FILES = {"pptx": _pptx_bytes(), "docx": _docx_bytes(),
         "xlsx": _xlsx_bytes(), "pdf": _pdf_bytes()}


def _row(doc_type: str) -> dict:
    # ⚠️ `public` + `published` and a DIFFERENT owner: the reader has to be someone other than
    # the owner, or `_notify_owner` correctly stays silent ("the owner is not told about their
    # own note") and the Inbox leg of this file would assert nothing.
    return dict(id=41, slug=f"card-{doc_type}", title=f"Q4 {doc_type} card", status="published",
                owner_email=DOC_OWNER, visibility="public", kind="document", html="",
                filters_json="", doc_type=doc_type, doc_object=f"obj-{doc_type}",
                doc_mime=reports.DOC_MIME[doc_type], doc_name=f"Q4.{doc_type}",
                doc_pages=len(FILES[doc_type]) and None, size_bytes=len(FILES[doc_type]),
                views=0, created_at=datetime(2026, 10, 1, tzinfo=timezone.utc),
                updated_at=datetime(2026, 10, 1, tzinfo=timezone.utc))


ROWS = {f"card-{kind}": _row(kind) for kind in FILES}


# ── the stub database ────────────────────────────────────────────────────────
# One fake for both modules: the rows these routes ask for are a handful of SELECTs, and
# splitting them across two fakes would only mean two places to update.

class _Cursor:
    def __init__(self, store: dict):
        self.store = store
        self._row = None
        self._rows = None
        self.rowcount = 0

    def execute(self, sql, params=None):
        text = " ".join(str(sql).split()).lower()
        params = tuple(params or ())
        self._row = None
        self._rows = None
        if "from ai_reports where slug" in text:
            self._row = ROWS.get(params[0])
        elif "from ai_report_anyone_links" in text:
            self._row = ROWS.get("card-docx")
        elif "update ai_reports set views" in text:
            self.rowcount = 1
        elif "from ai_report_colleague_shares" in text:
            self._row = None
        elif "from doc_annotations" in text and "order by" in text:
            # Two spellings of the same list: the signed-in one filters on (asset_type, slug),
            # the token one has `asset_type = 'report'` inlined. Both end with the slug.
            self._rows = sorted(
                [n for n in self.store["notes"].values() if n["asset_slug"] == params[-1]],
                key=lambda n: (n["page"], n["id"]))
        elif "select count(*) as note_count" in text:
            self._row = {"note_count": len([n for n in self.store["notes"].values()
                                            if n["asset_type"] == params[0]
                                            and n["asset_slug"] == params[1]])}
        elif "insert into doc_annotations" in text:
            note_id = self.store["seq"] = self.store["seq"] + 1
            now = datetime(2026, 10, 1, 9, 30, tzinfo=timezone.utc)
            asset_type, slug, asset_id, body, x, y, page, author = params
            note = {"id": note_id, "asset_type": asset_type, "asset_slug": slug,
                    "asset_id": asset_id, "body": body, "x_pct": x, "y_pct": y, "page": page,
                    "author_email": author, "resolved_at": None, "resolved_by": None,
                    "created_at": now, "updated_at": now}
            self.store["notes"][note_id] = note
            self._row = dict(note)
        elif "from doc_annotations where id = %s" in text:
            self._row = self.store["notes"].get(params[0])
        elif "update doc_annotations set body" in text:
            self.rowcount = 1
        elif "update doc_annotations set resolved_at" in text:
            self.rowcount = 1
        elif "delete from doc_annotations" in text:
            self.rowcount = 1
        else:
            raise AssertionError(f"the stub does not know this SQL: {text[:120]}")

    def fetchone(self):
        return self._row

    def fetchall(self):
        return self._rows or []

    def __enter__(self):
        return self

    def __exit__(self, *exc):
        return False


class _Conn:
    def __init__(self, store):
        self.store = store

    # `reports` opens its connections as `with _db() as conn:` and `annotations` calls
    # `.close()` by hand; the stub has to answer both shapes.
    def __enter__(self):
        return self

    def __exit__(self, *exc):
        return False

    def cursor(self, *a, **kw):
        return _Cursor(self.store)

    def commit(self):
        pass

    def rollback(self):
        pass

    def close(self):
        pass


STORE = {"notes": {}, "seq": 0}
EMITTED: list[dict] = []
SEEN: list[tuple] = []


# ── the app ──────────────────────────────────────────────────────────────────

def build_app() -> FastAPI:
    app = FastAPI()
    conn = lambda: _Conn(STORE)  # noqa: E731
    reports._ensure_table = lambda: None
    reports._db = conn
    # The guest link points at the card the note is on, or "does the guest see the notes"
    # would be answered by an empty document.
    reports._resolve_anyone_link = lambda token, request=None: ROWS["card-docx"]
    annotations._db = conn
    annotations.ensure_tables = lambda: None
    inbox.ensure_tables = lambda: None
    inbox.emit = lambda **kw: EMITTED.append(kw)
    inbox.forget_note = lambda note_id: None
    # The stub hands back the same bytes for every card, so the object key maps to a type.
    oss_storage.get_object = lambda bucket, key: FILES[key.split("-")[-1]]

    @app.middleware("http")
    async def _identity(request: Request, call_next):
        request.state.current_user = {"email": OWNER, "display_name": "Zhen Huang"}
        request.state.auth_kind = "browser"
        if request.method in ("POST", "PATCH", "DELETE"):
            body = await request.body()
            SEEN.append((request.method, request.url.path, body.decode("utf-8", "replace")))
        return await call_next(request)

    @app.get("/api/auth/me")
    async def _me():
        return {"authenticated": True, "enabled": False,
                "user": {"email": OWNER, "display_name": "Zhen Huang"}}

    @app.get("/r/{slug}", include_in_schema=False)
    async def _standalone(slug: str, request: Request, download: bool = False):
        # ⚠️ `download` must be passed explicitly — see the file header.
        return await reports.render_report(slug, request, download=False)

    @app.get("/s/{token}", include_in_schema=False)
    async def _guest(token: str, request: Request):
        return reports._document_page_response(reports._resolve_anyone_link(token), request,
                                               guest_token=token)

    @app.get("/s/{token}/document", include_in_schema=False)
    async def _guest_document(token: str, preview: str = Query("")):
        return await reports.anyone_link_document(token, preview)

    @app.get("/s/{token}/annotations", include_in_schema=False)
    async def _guest_notes(token: str):
        return await reports.anyone_link_annotations(token)

    app.include_router(reports.router, prefix="/api/reports")
    app.include_router(annotations_router.router, prefix="/api/annotations")

    @app.api_route("/api/{rest:path}", methods=["GET", "POST", "PATCH", "DELETE"],
                   include_in_schema=False)
    async def _anything(rest: str):
        SEEN.append(("unstubbed", "/api/" + rest, ""))
        return JSONResponse({})

    index = FRONTEND / "index.html"

    @app.get("/demo", include_in_schema=False)
    async def _demo():
        """A four-link index for `--serve`: this harness has no card wall, and clicking a
        card is how anyone actually judges whether a preview looks like Office."""
        links = "".join(
            f'<li><a href="/r/card-{kind}">{kind}</a> — '
            f'<a href="/s/{TOKEN}">as a guest link</a></li>' for kind in FILES)
        return Response(
            '<!DOCTYPE html><html><head><meta charset="utf-8"><title>Office previews</title>'
            '<style>body{font:16px/1.7 -apple-system,"PingFang SC",sans-serif;padding:40px;'
            'background:#f6f8fa}a{color:#1d6fd6}</style></head><body>'
            "<h1>Office 预览 + 备注 demo</h1>"
            "<p>每张卡片打开后：预览由读者浏览器渲染；在纸上右键可以留备注；"
            "访客链接只读、能看备注不能写。</p><ul>" + links + "</ul></body></html>",
            media_type="text/html")

    @app.get("/", include_in_schema=False)
    async def _index():
        return FileResponse(index, headers={"Cache-Control": "no-store"})

    app.mount("/", StaticFiles(directory=str(FRONTEND), html=True), name="static")
    return app


def check(name, ok, detail=""):
    print(("PASS  " if ok else "FAIL  ") + name + (("  — " + str(detail)) if detail else ""))
    if not ok:
        FAILURES.append(name)


def preview_frame(page, needle="preview"):
    """The preview document inside the reader page.

    Found by URL rather than by index: `page.frames[1]` is a guess about how many frames
    exist and in what order, and a guess that silently becomes the OUTER page is how a
    probe ends up measuring the reader chrome instead of the document. The needle is
    `preview`, not `preview.html`, because a guest link asks for the same document as
    `?preview=html`.
    """
    for candidate in page.frames:
        if needle in candidate.url:
            return candidate
    return None


# ── what the browser reports ─────────────────────────────────────────────────

# Inside a preview frame: the desk, the sheets, and the per-type chrome.
FRAME_PROBE = """
(() => {
  const sheets = [...document.querySelectorAll('.pv-page, .pv-slide, .pv-sheet')];
  const cs = getComputedStyle;
  return {
    bodyBg: cs(document.body).backgroundColor,
    sheets: sheets.length,
    shape: sheets.map(s => [Math.round(s.getBoundingClientRect().width),
                            Math.round(s.getBoundingClientRect().height)]),
    white: sheets.every(s => cs(s).backgroundColor === 'rgb(255, 255, 255)'),
    shadows: sheets.every(s => cs(s).boxShadow !== 'none'),
    captions: [...document.querySelectorAll('.pv-caption')].map(e => e.textContent.trim()),
    letters: [...document.querySelectorAll('.pv-colhead')].map(e => e.textContent.trim()),
    gutters: [...document.querySelectorAll('.pv-gutter')].map(e => e.textContent.trim()),
    tabs: [...document.querySelectorAll('.pv-tab')].map(e => e.textContent.trim()),
    images: document.querySelectorAll('.pv-page--pdf img').length,
    fontSize: sheets.length ? cs(sheets[0]).fontSize : null,
    noteLayer: !!document.querySelector('.doc-anno-layer--deck'),
    noteToggle: !!document.querySelector('.doc-anno-toggle'),
    remixicon: !!document.querySelector('link[href*="remixicon"]'),
  };
})()
"""


def main() -> int:
    import socket
    with socket.socket(socket.AF_INET6, socket.SOCK_STREAM) as probe:
        probe.settimeout(0.2)
        if probe.connect_ex(("::1", PORT)) == 0:
            print(f"port {PORT} is already serving something — stop it "
                  "(`pkill -f verify_office_notes_ui`) and re-run.")
            return 2

    server = uvicorn.Server(uvicorn.Config(build_app(), host="::1", port=PORT,
                                           log_level="warning"))
    threading.Thread(target=server.run, daemon=True).start()
    for _ in range(100):
        if getattr(server, "started", False):
            break
        time.sleep(0.05)
    base = f"http://[::1]:{PORT}"

    with sync_playwright() as playwright:
        browser = playwright.chromium.launch(headless=True,
                                             args=["--no-sandbox", "--no-proxy-server"])
        page = browser.new_page(viewport={"width": 1500, "height": 950})
        errors: list[str] = []
        page.on("pageerror", lambda exc: errors.append(str(exc)))

        # ── 1. every type looks like its own application ─────────────────────
        for kind in ("pptx", "docx", "xlsx", "pdf"):
            page.goto(f"{base}/r/card-{kind}", wait_until="domcontentloaded")
            page.wait_for_timeout(500)
            reader = page.evaluate(
                "() => ({frame: (document.getElementById('rpt-frame')||{}).src || '',"
                " src: (document.querySelector('iframe.dp-frame')||{}).src || '',"
                " hint: (document.querySelector('.dp-hint')||{}).textContent || ''})")
            check(f"{kind}: the reader page frames the preview",
                  reader["src"].endswith(f"api/reports/card-{kind}/document/preview.html"),
                  reader["src"])

            frame = preview_frame(page)
            probe = frame.evaluate(FRAME_PROBE) if frame else {}
            check(f"{kind}: {probe.get('sheets')} sheets on a grey desk",
                  probe.get("bodyBg") == "rgb(233, 237, 242)" and probe.get("sheets", 0) >= 1,
                  f"{probe.get('bodyBg')} / {probe.get('sheets')}")
            check(f"{kind}: the sheets are white and lifted off the desk",
                  bool(probe.get("white")) and bool(probe.get("shadows")), probe.get("shape"))
            check(f"{kind}: the preview carries the note layer",
                  bool(probe.get("noteLayer")) and bool(probe.get("noteToggle")),
                  probe.get("noteLayer"))
            check(f"{kind}: the note buttons have an icon font to draw from",
                  bool(probe.get("remixicon")))
            if kind == "pptx":
                check("pptx: a card per slide, each captioned",
                      probe.get("sheets") == 3 and probe.get("captions", [])[:3]
                      == ["Slide 1", "Slide 2", "Slide 3"], probe.get("captions"))
                check("pptx: the slide is laid out on a 16:9 canvas",
                      (probe.get("shape") or [[]])[0] == [1280, 720], probe.get("shape"))
            if kind == "docx":
                check("docx: the file's own page break becomes two sheets of paper",
                      probe.get("sheets") == 2
                      and probe.get("captions") == ["Page 1", "Page 2"], probe.get("captions"))
                check("docx: a page is paper-shaped, not a full-width column",
                      (probe.get("shape") or [[0, 0]])[0][0] == 820
                      and (probe.get("shape") or [[0, 0]])[0][1] >= 1123,
                      probe.get("shape"))
            if kind == "xlsx":
                check("xlsx: Excel's column letters are above the grid",
                      probe.get("letters", [])[:3] == ["A", "B", "C"], probe.get("letters"))
                check("xlsx: row 1 is the header row, so data starts at 2",
                      probe.get("gutters", [])[:3] == ["", "1", "2"], probe.get("gutters"))
                check("xlsx: the sheet names ride in a tab strip",
                      probe.get("tabs") == ["渠道口径"], probe.get("tabs"))
            if kind == "pdf":
                check("pdf: every page is a rasterised page image",
                      probe.get("images") == 2 and probe.get("sheets") == 2, probe.get("shape"))
        check("no page errors while rendering the four previews", not errors, errors[:2])

        # ── 2. a note belongs to the sheet it was left on ────────────────────
        page.goto(f"{base}/r/card-docx", wait_until="domcontentloaded")
        page.wait_for_timeout(600)
        frame = preview_frame(page)
        # ⚠️ `elementFromPoint` only answers for points INSIDE the viewport, and sheet 2 sits
        # below a 1150px-tall sheet 1 — so scroll it into view first, and give the note runtime
        # its scroll-sync beat before clicking, or the note is filed against the sheet the
        # reader was looking at a moment ago.
        frame.evaluate("() => document.querySelectorAll('.pv-page')[1]"
                       ".scrollIntoView({block: 'center'})")
        page.wait_for_timeout(1300)
        frame.evaluate("""() => {
            const box = document.querySelectorAll('.pv-page')[1].getBoundingClientRect();
            const x = box.left + box.width * 0.4;
            const y = Math.min(box.bottom - 20, Math.max(box.top + 20, window.innerHeight * 0.45));
            const target = document.elementFromPoint(x, y);
            target.dispatchEvent(new MouseEvent('contextmenu', {clientX: x, clientY: y,
                                                                bubbles: true, cancelable: true,
                                                                view: window}));
        }""")
        page.wait_for_timeout(250)
        menu = frame.evaluate("() => (document.querySelector('.doc-anno-menu')||{}).innerText || ''")
        check("right-clicking sheet 2 offers a note, and says which sheet",
              "Add a note here" in menu and "Page 2" in menu, repr(menu))
        frame.evaluate("() => document.querySelector('.doc-anno-menu button').click()")
        page.wait_for_timeout(250)
        frame.evaluate("""() => {
            const area = document.querySelector('.doc-anno-compose textarea');
            area.value = '这里的责任人写错了';
            document.querySelector('.doc-anno-compose .primary').click();
        }""")
        page.wait_for_timeout(700)

        posted = [body for method, path, body in SEEN if path.endswith("/api/annotations")]
        payload = json.loads(posted[-1]) if posted else {}
        check("the note was saved against the document card",
              payload.get("asset_type") == "report" and payload.get("slug") == "card-docx",
              payload)
        check("the note carries the sheet index, not a whole-scroll offset",
              payload.get("page") == 2, payload.get("page"))
        check("the note carries where on the sheet it was pinned",
              0 < float(payload.get("x", 0)) < 100 and 0 < float(payload.get("y", 0)) < 100,
              (payload.get("x"), payload.get("y")))

        bubble = frame.evaluate("""() => {
            const n = document.querySelector('.doc-anno');
            if (!n) return null;
            const box = n.getBoundingClientRect();
            const sheet = document.querySelectorAll('.pv-page')[1].getBoundingClientRect();
            return {inside: box.top >= sheet.top - 2 && box.bottom <= sheet.bottom + 2,
                    text: n.innerText, left: Math.round(box.left), top: Math.round(box.top)};
        }""")
        check("the bubble is drawn on sheet 2, where it was left",
              bool(bubble) and bubble["inside"] and "这里的责任人写错了" in bubble["text"], bubble)

        # Notes are scoped to the sheet in view (the same contract the deck already had:
        # "N note(s) on this page"). Scrolling to sheet 1 must therefore take the bubble
        # off screen, and scrolling back must bring it back where it was left.
        frame.evaluate("() => window.scrollTo(0, 0)")
        page.wait_for_timeout(1400)
        away = frame.evaluate("""() => ({
            bubbles: document.querySelectorAll('.doc-anno').length,
            count: (document.querySelector('.doc-anno-toggle-n')||{}).textContent || '',
        })""")
        check("a note is only drawn on the sheet it belongs to",
              away["bubbles"] == 0 and away["count"] == "0", away)
        frame.evaluate("() => document.querySelectorAll('.pv-page')[1]"
                       ".scrollIntoView({block: 'center'})")
        page.wait_for_timeout(1300)
        back = frame.evaluate("""() => {
            const n = document.querySelector('.doc-anno');
            if (!n) return null;
            const box = n.getBoundingClientRect();
            const sheet = document.querySelectorAll('.pv-page')[1].getBoundingClientRect();
            return {inside: box.top >= sheet.top - 2 && box.bottom <= sheet.bottom + 2,
                    count: (document.querySelector('.doc-anno-toggle-n')||{}).textContent || ''};
        }""")
        check("the bubble comes back on its own sheet, in the same place",
              bool(back) and back["inside"] and back["count"] == "1", back)

        # ── 3. the owner hears about it, with a name and a time ──────────────
        note_line = next((row for row in EMITTED if row.get("kind") == "note"), {}) or {}
        check("the note reaches the Inbox as a note event",
              bool(note_line) and note_line.get("recipient_email") == DOC_OWNER
              and note_line.get("note_id"), note_line)
        check("the Inbox line names the actor, the document and the sheet",
              note_line.get("actor_email") == OWNER
              and "Q4 docx card" in str(note_line.get("summary"))
              and "第 2 页" in str(note_line.get("summary")), note_line)
        # The sheet number the reader SAW has to be the number everyone else is told: the menu
        # said "Page 2", so the Inbox and the agent's one-liner must not say 3.
        check("the sheet number is the one the reader was shown",
              note_line.get("summary", "").endswith("这里的责任人写错了")
              and "第 3 页" not in str(note_line.get("summary")), note_line.get("summary"))
        notes = page.evaluate(
            f"() => fetch('{base}/api/annotations?type=report&slug=card-docx').then(r => r.json())")
        first = (notes.get("notes") or [{}])[0]
        check("the API hands an agent the author and the timestamp",
              first.get("author") == OWNER and bool(first.get("created_at"))
              and first.get("page") == 2, first)
        check("the agent-facing one-liner carries both too",
              OWNER in annotations.note_to_agent_line(first)
              and "page 2" in annotations.note_to_agent_line(first),
              annotations.note_to_agent_line(first))

        # ── 4. a guest may read the notes and never write one ────────────────
        page.goto(f"{base}/s/{TOKEN}", wait_until="domcontentloaded")
        page.wait_for_timeout(600)
        guest = preview_frame(page)
        check("a guest link gets the preview itself", bool(guest))
        # The note is on sheet 2, so the guest has to be looking at sheet 2 for it to be drawn —
        # notes are per sheet for everyone, signed in or not.
        guest.evaluate("() => document.querySelectorAll('.pv-page')[1]"
                       ".scrollIntoView({block: 'center'})")
        page.wait_for_timeout(1300)
        state = guest.evaluate("""() => ({
            layer: !!document.querySelector('.doc-anno-layer--deck'),
            notes: document.querySelectorAll('.doc-anno').length,
            title: (document.querySelector('.doc-anno-toggle')||{}).title || '',
            text: (document.querySelector('.doc-anno')||{}).innerText || '',
        })""")
        # The bubble names the author the way a colleague reads it ("reader", not the full
        # address) and stamps it with the time — the same two things the Inbox row carries.
        check("a guest link shows the notes already on the document, with author and time",
              state["layer"] and state["notes"] == 1
              and state["text"].startswith("reader") and "10-01" in state["text"]
              and "这里的责任人写错了" in state["text"], state)
        check("the guest's notes came back from the token route, not from a failed fetch",
              "could not be loaded" not in state["title"], state["title"])
        guest.evaluate("""() => {
            const el = document.elementFromPoint(400, 300);
            el.dispatchEvent(new MouseEvent('contextmenu', {clientX: 400, clientY: 300,
                                                            bubbles: true, cancelable: true,
                                                            view: window}));
        }""")
        page.wait_for_timeout(250)
        check("a guest is never offered a write route",
              not guest.evaluate("() => !!document.querySelector('.doc-anno-menu')"))

        browser.close()

    print()
    if FAILURES:
        print(f"{len(FAILURES)} FAILED: " + "; ".join(FAILURES))
        return 1
    print("all checks passed")
    return 0


def serve() -> int:
    """Leave the harness up so a human can look at the four previews."""
    uvicorn.run(build_app(), host="::1", port=SERVE_PORT, log_level="warning")
    return 0


if __name__ == "__main__":
    if "--serve" in sys.argv:
        raise SystemExit(serve())
    raise SystemExit(main())